# -*- coding: utf-8 -*-
"""
Full report generator for SCM NMMD Input/Output Alignment.
Outputs: BAO_CAO_NGHIEP_VU_DAU_VAO_DAU_RA_SCM_NMMD.docx
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left=None, right=None, top_sz="4", bottom_sz="4", left_sz="0", right_sz="0"):
    tcPr = cell._tc.get_or_add_tcPr()
    left_xml = f'<w:left w:val="single" w:sz="{left_sz}" w:space="0" w:color="{left}"/>' if left else '<w:left w:val="none"/>'
    right_xml = f'<w:right w:val="single" w:sz="{right_sz}" w:space="0" w:color="{right}"/>' if right else '<w:right w:val="none"/>'
    top_xml = f'<w:top w:val="single" w:sz="{top_sz}" w:space="0" w:color="{top}"/>' if top else '<w:top w:val="none"/>'
    bottom_xml = f'<w:bottom w:val="single" w:sz="{bottom_sz}" w:space="0" w:color="{bottom}"/>' if bottom else '<w:bottom w:val="none"/>'
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>{top_xml}{left_xml}{bottom_xml}{right_xml}</w:tcBorders>')
    tcPr.append(borders)

def make_callout(doc, text_list, title="LƯU Ý NGHIỆP VỤ QUAN TRỌNG", box_type="info"):
    border_color = "00552E" if box_type == "info" else ("D97706" if box_type == "warning" else "0284C7")
    bg_color = "F0FDF4" if box_type == "info" else ("FFFBEB" if box_type == "warning" else "F0F9FF")
    title_color = RGBColor(0, 85, 46) if box_type == "info" else (RGBColor(180, 83, 9) if box_type == "warning" else RGBColor(2, 132, 199))
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(10)
    run_t.font.bold = True
    run_t.font.color.rgb = title_color
    
    for item in text_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        run_item = p_item.add_run(f"• {item}")
        run_item.font.name = 'Calibri'
        run_item.font.size = Pt(9.5)
        run_item.font.color.rgb = RGBColor(51, 65, 85)
        
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_image_with_caption(doc, img_path, caption_text, figure_num):
    if not os.path.exists(img_path):
        p_err = doc.add_paragraph(f"[Không tìm thấy tệp hình ảnh: {img_path}]")
        p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(6.2))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap_lbl = p_cap.add_run(f"Hình {figure_num}: ")
    run_cap_lbl.font.name = 'Calibri'
    run_cap_lbl.font.size = Pt(9)
    run_cap_lbl.font.bold = True
    run_cap_lbl.font.color.rgb = RGBColor(0, 85, 46)
    
    run_cap_text = p_cap.add_run(caption_text)
    run_cap_text.font.name = 'Calibri'
    run_cap_text.font.size = Pt(9)
    run_cap_text.font.italic = True
    run_cap_text.font.color.rgb = RGBColor(100, 116, 139)

def style_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 85, 46)
    return p

def style_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def style_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(13, 130, 75)
    return p

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def format_custom_table(table, col_widths, headers, data, align_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "00552E")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Make header repeat on every page
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Format Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        r_trPr = table.rows[row_idx + 1]._tr.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        
        for col_idx, cell_value in enumerate(row_data):
            cell = row_cells[col_idx]
            cell.text = cell_value
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            set_cell_borders(cell, top="E2E8F0", bottom="E2E8F0")
            
            p = cell.paragraphs[0]
            if align_cols and col_idx in align_cols:
                p.alignment = align_cols[col_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(8.5)
                if "✔" in cell_value or "100%" in cell_value:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0, 85, 46)
                elif "ĐẦU VÀO" in cell_value or "ĐẦU RA" in cell_value:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(30, 41, 59)
                elif "NGHIÊM TRỌNG" in cell_value or "LỰA CHỌN" in cell_value:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(180, 83, 9)
                else:
                    r.font.color.rgb = RGBColor(51, 65, 85)

    # Set column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

print("Starting generation of document...")
