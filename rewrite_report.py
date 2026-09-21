# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 80, 40) # Emerald Dark Green
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate Dark
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_body_paragraph(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.first_line_indent = Pt(18)
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Times New Roman'
        r_bold.font.size = Pt(12)
        r_bold.font.bold = True
        
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = italic
    return p

def add_bullet_point(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Pt(18)
    
    r_bullet = p.add_run("• ")
    r_bullet.font.name = 'Times New Roman'
    r_bullet.font.size = Pt(12)
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = RGBColor(0, 80, 40)
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Times New Roman'
        r_bold.font.size = Pt(12)
        r_bold.font.bold = True
        
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Pt(14)
    p.paragraph_format.right_indent = Pt(14)
    
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def main():
    doc_path = r'D:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'
    doc = docx.Document(doc_path)
    
    print(f"Initial paragraph count: {len(doc.paragraphs)}")
    
    # Boundary: Keep paragraphs 0 to 227 intact. Delete paragraphs from 228 to end.
    # In python-docx, to delete paragraphs safely:
    body_element = doc._body._element
    paragraphs = doc.paragraphs
    
    if len(paragraphs) > 228:
        for p in paragraphs[228:]:
            p._element.getparent().remove(p._element)
            
    print(f"Paragraph count after truncation: {len(doc.paragraphs)}")
    
    # Now append the complete rewritten content starting from CHƯƠNG 1
    
    # =========================================================================
    # CHƯƠNG 1
    # =========================================================================
    add_heading_1(doc, "CHƯƠNG 1: TỔNG QUAN VỀ ĐƠN VỊ THỰC TẬP VÀ ĐỀ TÀI DỰ ÁN")
    
    add_heading_2(doc, "1.1. Giới thiệu tổng quan về Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II)")
    add_body_paragraph(doc, "Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II) là cụm nhà máy sản xuất da giày thể thao xuất khẩu chủ lực trực thuộc Tập đoàn TBS Group (Thái Bình Corporation) — một trong những tập đoàn sản xuất da giày, túi xách và hạ tầng logistics hàng đầu tại Việt Nam.")
    add_bullet_point(doc, "Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2", "Tên đầy đủ: ")
    add_bullet_point(doc, "Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II Workspace)", "Tên viết tắt / Thương hiệu: ")
    add_bullet_point(doc, "2/434, Khu phố Bình Đáng, Phường Bình Hòa, Thành phố Thuận An, Tỉnh Bình Dương", "Địa chỉ trụ sở chính & Cụm Nhà máy: ")
    add_bullet_point(doc, "3700148644", "Mã số thuế: ")
    add_bullet_point(doc, "(0274) 3758 888 / (0297) 3922 888", "Điện thoại văn phòng: ")
    add_bullet_point(doc, "https://www.tbsgroup.vn/", "Website chính thức: ")
    
    add_body_paragraph(doc, "Lĩnh vực hoạt động sản xuất cốt lõi của Nhà Máy 2 bao gồm:")
    add_bullet_point(doc, "Là đơn vị sản xuất chủ lực chuyên hoàn thiện các dòng sản phẩm giày thể thao chất lượng cao cho Chuỗi Skechers toàn cầu xuất khẩu sang thị trường Mỹ, Châu Âu và quốc tế.", "1. Sản xuất & Gia công Da Giày Thể thao Xuất khẩu: ")
    add_bullet_point(doc, "Ép đế PU, lưu hóa đế cao su và thiết kế khuôn mẫu phụ liệu cung ứng trực tiếp cho các dây chuyền gò, may, ráp tại nhà máy.", "2. Sản xuất Phụ liệu & Đế PU / Cao su: ")

    add_heading_2(doc, "1.2. Cơ cấu tổ chức và chức năng nhiệm vụ của Nhà Máy 2 (TBS II)")
    add_body_paragraph(doc, "Trên cơ sở định hướng phát triển của Tập đoàn TBS Group, Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 thực hiện các chức năng chính bao gồm:")
    add_bullet_point(doc, "Vận hành 33 dây chuyền may, gò, ráp và đóng gói sản phẩm giày thể thao xuất khẩu đạt chuẩn chất lượng quốc tế.", "- Khối Sản xuất & Gia công Da Giày Skechers: ")
    add_bullet_point(doc, "Thực hiện kiểm định Gemba Walk 5S định kỳ tại khu vực sản xuất, quản lý các thẻ lỗi Andon và đảm bảo tiêu chuẩn chất lượng sản phẩm.", "- Khối Quản lý Chất lượng (QC) & Gemba Walk: ")
    add_bullet_point(doc, "Tiếp nhận, đánh giá, xếp hàng duyệt (IE Queue) và đưa vào ứng dụng các đề xuất cải tiến Kaizen nhằm nâng cao năng suất chuyền.", "- Khối Cải tiến liên tục (CI/IE - Kaizen Engine): ")
    add_bullet_point(doc, "Quản lý lịch trình, phân công nhiệm vụ cho Kỹ sư IE và CBCNV trên hệ thống thẻ Kanban điện tử.", "- Quản lý Thẻ Công việc Kanban: ")
    add_bullet_point(doc, "Bà Dư Thị Thanh Tình (Trưởng phòng Hành chính & Nhân sự).", "- Cán bộ hướng dẫn tại đơn vị: ")

    add_heading_2(doc, "1.3. Tổng quan về đề tài: Xây dựng Hệ thống Quản trị số hóa Vận hành Văn Phòng Chuỗi Skechers - TBS Group")
    add_body_paragraph(doc, "Trong bối cảnh chuyển đổi số công nghiệp 4.0 và yêu cầu tối ưu hóa quy trình vận hành chuỗi cung ứng sản xuất da giày, việc xây dựng một giải pháp quản trị số hóa tập trung cho Văn Phòng Chuỗi Skechers là vô cùng cấp thiết. Đề tài tập trung xây dựng hệ thống TBS II Workspace với các mục tiêu trọng tâm:")
    add_bullet_point(doc, "Loại bỏ quy trình báo cáo giấy thủ công, chuyển sang giao diện điều hành số 24/7.", "(1) Số hóa quản trị vận hành: ")
    add_bullet_point(doc, "Xây dựng bảng Kanban kéo thả, cập nhật trạng thái realtime (TODO, IN_PROGRESS, REVIEW, DONE), bắt buộc nhập nội dung kết quả thực hiện khi hoàn thành.", "(2) Tự động hóa quản lý thẻ công việc: ")
    add_bullet_point(doc, "Chuẩn hóa quy trình đăng ký Kaizen, xếp hàng duyệt IE Queue, kiểm định Gemba 5S và cảnh báo sự cố Andon.", "(3) Chuẩn hóa quy trình CI/Kaizen & Gemba Walk: ")
    add_bullet_point(doc, "Phát triển cơ chế lưu vết thao tác người dùng (Audit Logs) gắn kèm Tên CBNV, kết hợp cơ chế sao lưu dữ liệu tự động lên Google Drive xuất định dạng JSON và Excel CSV UTF-8 BOM.", "(4) Giám sát an toàn & Sao lưu realtime: ")

    add_heading_2(doc, "1.4. Kế hoạch và tiến độ thực tập (11 tuần từ 17/08/2026 đến 31/10/2026)")
    add_body_paragraph(doc, "Quá trình thực tập tại Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II) được triển khai theo kế hoạch 11 tuần dưới sự giám sát trực tiếp của cán bộ hướng dẫn:")
    add_bullet_point(doc, "Khảo sát quy trình nghiệp vụ thực tế tại Văn phòng Chuỗi Skechers TBS Group, phân tích nhu cầu số hóa và lập kế hoạch thực tập.", "• Tuần 1 (17/08 – 23/08/2026): ")
    add_bullet_point(doc, "Xây dựng kiến trúc Serverless Edge Backend dựa trên Cloudflare Workers và cơ sở dữ liệu Cloudflare D1.", "• Tuần 2 (24/08 – 30/08/2026): ")
    add_bullet_point(doc, "Phát triển giao diện người dùng Next.js 14 App Router, tích hợp Tailwind CSS và Tabler Icons.", "• Tuần 3 (31/08 – 06/09/2026): ")
    add_bullet_point(doc, "Xây dựng phân hệ Quản lý Thẻ Công việc Kanban (/work/my-tasks, /work/tasks), xử lý sự kiện Drag & Drop và Modal nhập kết quả result_description.", "• Tuần 4 (07/09 – 13/09/2026): ")
    add_bullet_point(doc, "Phát triển phân hệ Cải tiến Kaizen IE và Kiểm định Gemba Walk (/work/cn-ci, /work/kaizen, /work/gemba).", "• Tuần 5–6 (14/09 – 27/09/2026): ")
    add_bullet_point(doc, "Xây dựng cơ chế Real-time Event-Driven Backup lên Google Drive qua Google Apps Script WebApp, hỗ trợ xuất định dạng JSON và Excel CSV UTF-8 BOM có cột Tên CBNV.", "• Tuần 7–8 (28/09 – 11/10/2026): ")
    add_bullet_point(doc, "Tối ưu hóa UI/UX chống AI Slop theo chuẩn Taste Skill, tích hợp hệ thống Ghi nhật ký Thao tác Audit Logs (/work/admin/audit-logs) và chuẩn bảo mật RFC 9116.", "• Tuần 9–10 (12/10 – 25/10/2026): ")
    add_bullet_point(doc, "Kiểm thử tổng thể, tối ưu hiệu năng Edge Computing, nghiệm thu hệ thống và hoàn thiện báo cáo thực tập.", "• Tuần 11 (26/10 – 31/10/2026): ")

    # =========================================================================
    # CHƯƠNG 2
    # =========================================================================
    add_heading_1(doc, "CHƯƠNG 2: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")

    add_heading_2(doc, "2.1. Yêu cầu bài toán và yêu cầu chức năng hệ thống")
    add_body_paragraph(doc, "Hệ thống Quản trị số hóa TBS II Workspace được thiết kế nhằm giải quyết toàn diện bài toán số hóa vận hành nhà máy sản xuất da giày Skechers với các nhóm yêu cầu chức năng chính:")
    add_bullet_point(doc, "Hỗ trợ đăng nhập đa vai trò (RBAC), bảo mật JWT Token, lưu thông tin phiên làm việc và hiển thị cá nhân hóa theo từng cán bộ.", "1. Phân hệ Quản lý Tài khoản & Phân quyền: ")
    add_bullet_point(doc, "Giao diện bảng Kanban trực quan, cho phép kéo thả chuyển trạng thái công việc, bắt buộc nhập mô tả kết quả hoàn thành khi chuyển sang DONE.", "2. Phân hệ Quản lý Thẻ Công việc Kanban: ")
    add_bullet_point(doc, "Tiếp nhận đề xuất cải tiến từ dây chuyền, xếp hàng thẩm định Kỹ sư IE (IE Queue), tính toán thời gian hoàn vốn và theo dõi trạng thái phê duyệt.", "3. Phân hệ Cải tiến Kaizen IE: ")
    add_bullet_point(doc, "Thực hiện chấm điểm 5S Gemba Walk tại 33 chuyền sản xuất, ghi nhận sự cố Andon và gửi thông báo xử lý tức thì.", "4. Phân hệ Kiểm định Gemba Walk & Andon: ")
    add_bullet_point(doc, "Đăng ký sử dụng phòng họp trực tuyến và đặt lịch xe công tác liên tỉnh cho cán bộ công nhân viên.", "5. Phân hệ Phòng họp & Chuyến xe Công tác: ")
    add_bullet_point(doc, "Tự động ghi nhận mọi truy cập/chỉnh sửa của người dùng, phân tích IP, User-Agent, gắn kèm Tên CBNV và kích hoạt sao lưu realtime lên Google Drive.", "6. Phân hệ Audit Logs & Real-time Backup: ")

    add_heading_2(doc, "2.2. Kiến trúc tổng thể hệ thống Edge Serverless (Cloudflare Stack)")
    add_body_paragraph(doc, "Hệ thống được thiết kế theo kiến trúc Serverless Edge Computing hiện đại, loại bỏ sự phụ thuộc vào máy chủ ảo truyền thống (VPS Express.js) để tối ưu hóa chi phí vận hành và đảm bảo tốc độ phản hồi < 50ms toàn cầu:")
    add_bullet_point(doc, "Xây dựng bằng Next.js 14.2.15 (Static Export 'out/'), React 18, Tailwind CSS và Tabler Icons. Giao diện được thiết kế theo định hướng Taste Skill (chống AI Slop), mang phong cách nhận diện thương hiệu TBS Group Emerald Green (#006838) sang trọng.", "• Presentation Layer (Frontend PWA): ")
    add_bullet_point(doc, "Vận hành hoàn toàn trên Cloudflare Workers Edge Network (file script 'public/_worker.js' & 'out/_worker.js'). Lớp này tiếp nhận các yêu cầu HTTP/HTTPS API, thực thi logic nghiệp vụ, xác thực JWT và tương tác trực tiếp với cơ sở dữ liệu.", "• Edge Business Logic Layer (Worker Handler): ")
    add_bullet_point(doc, "Sử dụng Cloudflare D1 Serverless Relational SQLite Database (Binding 'env.DB' với cơ sở dữ liệu 'vpchuoiskechers-db'), lưu trữ toàn bộ dữ liệu bảng công việc, nhật ký audit, Kaizen, Gemba và tài khoản người dùng.", "• Database Layer (Cloudflare D1): ")

    add_heading_2(doc, "2.3. Cơ chế Sao lưu Tự động Realtime Google Drive (JSON & Excel CSV)")
    add_body_paragraph(doc, "Một điểm nhấn kiến trúc quan trọng của hệ thống là cơ chế Sao lưu sự kiện thời gian thực (Real-time Event-Driven Backup) kết nối trực tiếp với Google Drive thông qua Google Apps Script WebApp endpoint (env.GDRIVE_WEBAPP_URL):")
    add_body_paragraph(doc, "Để giải quyết triệt để lỗi 403 storageQuotaExceeded (do tài khoản Google Service Account có dung lượng lưu trữ 0 GB nên bị Google Drive API chặn khi tạo file trong thư mục cá nhân), hệ thống đã chuyển sang mô hình kết nối Google Apps Script WebApp. WebApp này được triển khai dưới quyền tài khoản chủ sở hữu tbsgroup2026@gmail.com (15 GB dung lượng khả dụng), nhận dữ liệu sao lưu POST từ Cloudflare Worker và ghi file trực tiếp vào thư mục gốc 'Văn Phòng Chuỗi'.")
    
    add_body_paragraph(doc, "Hệ thống tự động phân loại và ghi dữ liệu sao lưu vào 6 thư mục chuyên biệt trên Google Drive:")
    add_bullet_point(doc, "Sao lưu toàn bộ 27 bảng dữ liệu hệ thống (sys_my_tasks, meeting_rooms, business_trips, audit_logs...)", "1. Thư mục 00_Tong_Hop_Full_Database: ")
    add_bullet_point(doc, "Sao lưu chi tiết lịch sử thao tác và truy cập của người dùng.", "2. Thư mục 01_Nhat_Ky_Thao_Tac_Audit_Logs: ")
    add_bullet_point(doc, "Sao lưu danh sách tài khoản, hồ sơ nhân sự và vai trò hệ thống.", "3. Thư mục 02_Tai_Khoan_Nguoi_Dung_Users: ")
    add_bullet_point(doc, "Sao lưu toàn bộ đề xuất cải tiến Kaizen và trạng thái thẩm định.", "4. Thư mục 03_Sang_Kien_Cai_Tien_Kaizen: ")
    add_bullet_point(doc, "Sao lưu hồ sơ đánh giá Gemba 5S và lịch sử sự cố Andon.", "5. Thư mục 04_Quan_Ly_Gemba_Andon: ")
    add_bullet_point(doc, "Sao lưu lịch đăng ký phòng họp và chuyến xe công tác.", "6. Thư mục 05_Dat_Phong_Hop_Rooms: ")

    add_body_paragraph(doc, "Mỗi sự kiện sao lưu tự động sinh ra song song 2 định dạng tệp tin:")
    add_bullet_point(doc, "Chứa cấu trúc JSON nguyên bản, dùng cho công tác khôi phục dữ liệu hệ thống (Database Recovery).", "• File .json: ")
    add_bullet_point(doc, "Tự động chuyển đổi dữ liệu mảng thành bảng tính CSV mã hóa UTF-8 BOM (\\uFEFF). Định dạng này giúp người dùng nhấp đúp là mở trực tiếp trên Microsoft Excel hoặc Google Sheets hiển thị chuẩn 100% Tiếng Việt có dấu, đồng thời tự động bổ sung cột 'Tên CBNV' ở vị trí đầu tiên.", "• File .csv (Excel Format): ")

    add_heading_2(doc, "2.4. Thiết kế Cơ sở Dữ liệu Cloudflare D1 (Database Schema)")
    add_body_paragraph(doc, "Cơ sở dữ liệu Cloudflare D1 (vpchuoiskechers-db) bao gồm 27 bảng dữ liệu được chuẩn hóa. Các bảng cốt lõi của hệ thống bao gồm:")
    
    add_bullet_point(doc, "Lưu trữ thông tin công việc Kanban (id, code, title, description, department_id, assignee_emp_code, assignee_name, status [TODO/IN_PROGRESS/REVIEW/DONE], progress, checklist, result_description, updated_at).", "1. Bảng sys_my_tasks: ")
    add_bullet_point(doc, "Lưu trữ nhật ký vết thao tác người dùng (id, user_id, emp_code, emp_name [Tên CBNV], role_code, module, action, record_id, changes_json, ip_address, user_agent, created_at).", "2. Bảng audit_logs: ")
    add_bullet_point(doc, "Lưu trữ đề xuất cải tiến Kaizen (id, code, title, proposer_emp_code, department, status, cost_saving, payback_period_months, created_at).", "3. Bảng ci_kaizen_proposals: ")
    add_bullet_point(doc, "Lưu trữ kết quả kiểm định 5S Gemba Walk (id, code, factory, line, auditor_emp_code, score, status, created_at).", "4. Bảng gemba_records: ")
    add_bullet_point(doc, "Lưu trữ lịch đặt phòng họp và xe công tác.", "5. Bảng room_bookings & business_trips: ")
    add_bullet_point(doc, "Lưu trữ danh sách nhân viên và tài khoản người dùng.", "6. Bảng sys_users & user_profile: ")

    add_heading_2(doc, "2.5. Cơ chế Phân quyền RBAC và Chuẩn Bảo mật RFC 9116")
    add_body_paragraph(doc, "Hệ thống áp dụng cơ chế Phân quyền theo Vai trò (Role-Based Access Control - RBAC) với 6 vai trò chính:")
    add_bullet_point(doc, "Quyền xem tổng quan chỉ số toàn chuỗi, duyệt các đề xuất chiến lược.", "• TONG_GIAM_DOC (TGĐ-001): ")
    add_bullet_point(doc, "Quyền quản lý khối vận hành, theo dõi tiến độ các nhà máy.", "• PHO_TONG_GIAM_DOC (PTGĐ-002): ")
    add_bullet_point(doc, "Quyền quản lý khối sản xuất & tổ hợp nhà máy.", "• GIAM_DOC (GĐ-003): ")
    add_bullet_point(doc, "Quyền quản lý chất lượng (QC) và Gemba 5S.", "• PHO_GIAM_DOC (PGĐ-004): ")
    add_bullet_point(doc, "Quyền thực hiện công việc cá nhân, nộp Kaizen, đặt phòng họp.", "• CBCNV (202608001 / 202608003): ")
    add_bullet_point(doc, "Toàn quyền quản trị hệ thống, cấu hình phân hệ và xem nhật ký audit logs.", "• SYSTEM_ADMIN (ADMIN-2026): ")
    
    add_body_paragraph(doc, "Về an ninh bảo mật, hệ thống tuân thủ chuẩn an ninh quốc tế RFC 9116 bằng việc cung cấp endpoint /.well-known/security.txt, tự động khử các dữ liệu nhạy cảm (password, token, secret -> [REDACTED]) trước khi đưa vào dữ liệu sao lưu.")

    # =========================================================================
    # CHƯƠNG 3
    # =========================================================================
    add_heading_1(doc, "CHƯƠNG 3: TRIỂN KHAI VÀ PHÁT TRIỂN CÁC PHÂN HỆ CHỨC NĂNG")

    add_heading_2(doc, "3.1. Phân hệ Trung tâm Điều hành Workspace Hub (/work)")
    add_body_paragraph(doc, "Trang chủ Workspace Hub (/work) đóng vai trò là bảng điều khiển trung tâm cho toàn bộ Văn Phòng Chuỗi Skechers. Giao diện được thiết kế lại hoàn toàn theo định hướng Taste Skill (chống AI Slop):")
    add_bullet_point(doc, "Phối màu gradient tối Obsidian Emerald (from-[#041a13] via-[#004d29] to-[#01140c]), họa tiết lưới chìm mờ và hiệu ứng ambient glow xanh ngọc bích.", "• Hero Welcome Card: ")
    add_bullet_point(doc, "Các thẻ phân hệ (Trang chủ, Nhân sự - Hành chính, Dự án của tôi) thiết kế theo dạng thẻ thủy tinh bo góc 16px, hiệu ứng hover viền phản quang và chuyển màu icon mượt mà.", "• Thẻ Phân hệ Nghiệp vụ: ")
    add_bullet_point(doc, "Cung cấp các phím tắt truy cập nhanh 24/7 đến các tính năng: Công việc, Quản lý công việc, Phòng họp, Công tác.", "• Thanh Truy cập Nhanh (Focus Quick Access): ")

    add_heading_2(doc, "3.2. Phân hệ Quản lý Thẻ Công việc Kanban (/work/my-tasks, /work/tasks)")
    add_body_paragraph(doc, "Phân hệ Kanban quản lý công việc được phát triển với khả năng tương tác cao:")
    add_bullet_point(doc, "Cho phép kéo thả các thẻ công việc giữa 4 cột trạng thái: Cần làm (TODO), Đang làm (IN_PROGRESS), Chờ duyệt (REVIEW), Hoàn thành (DONE).", "• Xử lý Drag & Drop mượt mà: ")
    add_bullet_point(doc, "Khắc phục triệt để hiện tượng giật lắc giao diện bằng việc sử dụng useRef (draggedTaskIdRef) và fallback window._draggedTaskId, kết hợp bỏ qua sự kiện dragLeave từ các phần tử con.", "• Tối ưu hóa sự kiện Kéo thả: ")
    add_bullet_point(doc, "Khi người dùng kéo thả thẻ công việc sang cột DONE, hệ thống tự động bật Modal bắt buộc người dùng nhập kết quả thực hiện (result_description). Nút 'Xác nhận hoàn thành' chỉ kích hoạt khi đã có nội dung mô tả.", "• Modal Mô tả Kết quả Hoàn thành: ")
    add_bullet_point(doc, "Áp dụng cơ chế Optimistic UI để cập nhật trạng thái thẻ tức thì trên màn hình người dùng, đồng thời gửi yêu cầu PATCH /api/tasks đến Cloudflare Worker để cập nhật cơ sở dữ liệu D1 và kích hoạt sao lưu Google Drive realtime.", "• Đồng bộ D1 realtime: ")

    add_heading_2(doc, "3.3. Phân hệ Cải tiến Kaizen IE và Kiểm định Gemba Walk (/work/cn-ci, /work/kaizen, /work/gemba)")
    add_body_paragraph(doc, "Phân hệ hỗ trợ số hóa toàn bộ quy trình cải tiến chất lượng và vận hành nhà máy:")
    add_bullet_point(doc, "Giao diện đăng ký sáng kiến cải tiến, tự động đưa vào hàng chờ thẩm định của Kỹ sư IE (IE Queue), hỗ trợ tính toán thời gian hoàn vốn.", "• Đăng ký & Thẩm định Kaizen: ")
    add_bullet_point(doc, "Số hóa bảng chấm điểm 5S Gemba Walk cho 33 dây chuyền sản xuất, ghi nhận hình ảnh minh chứng và phát tín hiệu cảnh báo sự cố Andon.", "• Gemba Walk 5S Audit & Andon: ")

    add_heading_2(doc, "3.4. Phân hệ Đăng ký Phòng họp & Chuyến xe Công tác (/rooms, /business-trip)")
    add_body_paragraph(doc, "Cung cấp công cụ tiện ích văn phòng trực tuyến:")
    add_bullet_point(doc, "Trực quan hóa lịch trống của các phòng họp, hỗ trợ đăng ký và duyệt đặt phòng trực tuyến.", "• Đặt phòng họp: ")
    add_bullet_point(doc, "Đăng ký nhu cầu di chuyển công tác liên tỉnh, sắp xếp lịch điều xe và quản lý chi phí chuyến đi.", "• Quản lý xe công tác: ")

    add_heading_2(doc, "3.5. Phân hệ Nhật ký Thao tác Audit Logs (/work/admin/audit-logs)")
    add_body_paragraph(doc, "Đảm bảo tính minh bạch và an toàn thông tin cho toàn hệ thống:")
    add_bullet_point(doc, "Mọi thao tác thêm, sửa, xóa, chuyển trạng thái công việc đều được hàm recordAuditLog ghi lại vào bảng audit_logs.", "• Tự động lưu vết thao tác: ")
    add_bullet_point(doc, "Dữ liệu log được tự động bổ sung cột emp_name ('Tên CBNV') thông qua bảng ánh xạ mã nhân viên, giúp người quản trị dễ dàng theo dõi cán bộ thực hiện thay vì chỉ nhìn thấy mã số kỹ thuật.", "• Bổ sung cột Tên CBNV: ")

    add_heading_2(doc, "3.6. Quy trình CI/CD và Triển khai Đám mây Cloudflare")
    add_body_paragraph(doc, "Quy trình đóng gói và triển khai ứng dụng được tự động hóa hoàn toàn:")
    add_bullet_point(doc, "Chạy lệnh npm run build để Tailwind CSS biên dịch file style public/compiled-tailwind.css, Next.js xuất ra thư mục tĩnh out/, và tự động sao chép public/_worker.js sang out/_worker.js.", "1. Đóng gói ứng dụng (Build Step): ")
    add_bullet_point(doc, "Sử dụng công cụ Cloudflare Wrangler (npx wrangler deploy) để đẩy toàn bộ trang tĩnh và Worker handler lên mạng lưới toàn cầu Cloudflare.", "2. Triển khai siêu tốc (Deployment Step): ")

    # =========================================================================
    # CHƯƠNG 4
    # =========================================================================
    add_heading_1(doc, "CHƯƠNG 4: KIỂM THỬ, ĐÁNH GIÁ VÀ KẾT LUẬN")

    add_heading_2(doc, "4.1. Kiểm thử chức năng và hiệu năng hệ thống")
    add_body_paragraph(doc, "Hệ thống đã trải qua quá trình kiểm thử toàn diện trên cả 159 routes giao diện và các endpoint API:")
    add_bullet_point(doc, "Kiểm thử thành công các thao tác đăng nhập, kéo thả Kanban, bật modal nhập result_description, nộp Kaizen, đặt phòng họp.", "• Kiểm thử Chức năng (Functional Testing): ")
    add_bullet_point(doc, "Nhờ kiến trúc Edge Serverless trên Cloudflare Workers, thời gian phản hồi API trung bình đạt < 50ms, thời gian khởi động Worker (Startup Time) chỉ từ 11 - 16ms.", "• Kiểm thử Hiệu năng (Performance Testing): ")

    add_heading_2(doc, "4.2. Kiểm thử cơ chế Sao lưu Google Drive (JSON & Excel CSV)")
    add_body_paragraph(doc, "Tiến hành gọi endpoint kiểm thử /api/test-backup để đánh giá khả năng sao lưu thời gian thực:")
    add_bullet_point(doc, "Tất cả 6 thư mục mục tiêu trên Google Drive (00_Tong_Hop_Full_Database, 01_Nhat_Ky_Thao_Tac_Audit_Logs...) đều nhận đầy đủ file .json và file .csv Excel.", "• Kết quả kiểm thử: ")
    add_bullet_point(doc, "File Excel (.csv) được mở trực tiếp bằng Microsoft Excel và Google Sheets, hiển thị chuẩn 100% Tiếng Việt có dấu nhờ mã hóa UTF-8 BOM, đồng thời cột 'Tên CBNV' hiển thị rõ ràng ở vị trí đầu tiên.", "• Đánh giá định dạng Excel: ")

    add_heading_2(doc, "4.3. Đánh giá kết quả đạt được và bài học kinh nghiệm")
    add_body_paragraph(doc, "Sau 11 tuần thực tập tại Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II), em đã hoàn thành 100% các mục tiêu đề ra:")
    add_bullet_point(doc, "Xây dựng và đưa vào vận hành thành công Hệ thống Quản trị số hóa Văn Phòng Chuỗi Skechers - TBS Group trên nền tảng Cloudflare Edge Serverless.", "1. Về mặt sản phẩm: ")
    add_bullet_point(doc, "Làm chủ các công nghệ hiện đại như Next.js 14 App Router, Cloudflare Workers, Cloudflare D1 Database, Google Apps Script WebApp integration, và kỹ thuật thiết kế UI Taste Skill.", "2. Về mặt kiến thức kỹ thuật: ")
    add_bullet_point(doc, "Rèn luyện tư duy lập trình chuyên nghiệp, kỹ năng giải quyết sự cố thực tế (như khắc phục lỗi Quota 403 Google Drive và giật lắc Drag & Drop), kỹ năng làm việc nhóm và quản lý tiến độ.", "3. Về mặt kỹ năng mềm: ")

    add_heading_2(doc, "4.4. Định hướng phát triển và đề xuất kiến nghị")
    add_bullet_point(doc, "Tiếp tục nâng cấp phân hệ AI Automation để hỗ trợ phân tích dự báo sự cố Andon và gợi ý giải pháp Kaizen tự động dựa trên dữ liệu lịch sử.", "• Định hướng phát triển hệ thống: ")
    add_bullet_point(doc, "Nhà trường nên tăng cường thời lượng thực hành các công nghệ đám mây Serverless (Cloudflare, AWS), kiến trúc Edge Computing và các kỹ thuật xây dựng sản phẩm thực tế trong doanh nghiệp.", "• Kiến nghị đối với chương trình đào tạo: ")

    # =========================================================================
    # ĐÁNH GIÁ THỰC TIỄN & KIẾN NGHỊ
    # =========================================================================
    add_heading_1(doc, "ĐÁNH GIÁ THỰC TIỄN VÀ MỘT SỐ KIẾN NGHỊ")
    
    add_heading_2(doc, "Một số đánh giá về hoạt động của đơn vị thực tập")
    add_body_paragraph(doc, "Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II) là một môi trường làm việc chuyên nghiệp, năng động và đi đầu trong việc ứng dụng công nghệ chuyển đổi số vào quản trị vận hành sản xuất da giày. Đơn vị đã tạo điều kiện thuận lợi về hạ tầng, tài liệu kỹ thuật và sự hướng dẫn tận tình từ các cán bộ quản lý giúp sinh viên hoàn thành tốt đề tài thực tập.")

    add_heading_2(doc, "Bài học kinh nghiệm tích lũy trong quá trình thực tập")
    add_body_paragraph(doc, "Quá trình thực tập mang lại nhiều bài học giá trị: (1) Tư duy thiết kế hệ thống phải luôn bám sát nhu cầu thực tế của người dùng doanh nghiệp; (2) Việc lựa chọn kiến trúc công nghệ phù hợp (như Serverless Edge Computing) giúp tiết kiệm đáng kể chi phí hạ tầng và tối ưu hiệu năng; (3) Kỹ năng giải quyết sự cố (troubleshooting) dựa trên bằng chứng log thực tế là yếu tố quyết định sự thành công của dự án phần mềm.")

    add_heading_2(doc, "Một số kiến nghị đối với Nhà trường và Doanh nghiệp")
    add_body_paragraph(doc, "Kiến nghị đối với Nhà trường: Cần cập nhật thêm các học phần về Điện toán đám mây Serverless, Điện toán biên (Edge Computing) và quy trình CI/CD tự động vào chương trình giảng dạy chính khóa. Đề xuất mở rộng các chương trình hợp tác doanh nghiệp để sinh viên có cơ hội tiếp cận sớm với các dự án thực tế.")
    add_body_paragraph(doc, "Kiến nghị đối với Doanh nghiệp: Tiếp tục mở rộng quy mô ứng dụng hệ thống TBS II Workspace sang các cụm nhà máy khác thuộc Tập đoàn TBS Group, đồng thời đầu tư thêm các thiết bị máy quét IoT để tự động hóa khâu điểm danh và ghi nhận dữ liệu Gemba Walk.")

    # =========================================================================
    # KẾT LUẬN
    # =========================================================================
    add_heading_1(doc, "KẾT LUẬN")
    add_body_paragraph(doc, "Báo cáo thực tập đã phản ánh toàn bộ quá trình 11 tuần làm việc và nghiên cứu của em tại Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II). Việc xây dựng thành công Hệ thống Quản trị số hóa Vận hành Văn Phòng Chuỗi Skechers - TBS Group (TBS II Workspace) không chỉ giải quyết hiệu quả bài toán số hóa vận hành nhà máy mà còn khẳng định tính đúng đắn của việc ứng dụng công nghệ Serverless Edge Computing vào môi trường doanh nghiệp. Em xin chân thành cảm ơn sự hướng dẫn tận tình của Quý Thầy Cô Viện Công nghệ số - Trường Đại học Thủ Dầu Một và Ban Lãnh đạo, cán bộ hướng dẫn tại Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 đã hỗ trợ em hoàn thành tốt đợt thực tập này.")

    # =========================================================================
    # TÀI LIỆU THAM KHẢO
    # =========================================================================
    add_heading_1(doc, "TÀI LIỆU THAM KHẢO")
    add_bullet_point(doc, "Next.js Documentation (2026). App Router and Static Export Guides. Retreived from https://nextjs.org/docs")
    add_bullet_point(doc, "Cloudflare Workers Documentation (2026). Edge Computing and D1 Database Reference. Retreived from https://developers.cloudflare.com/workers/")
    add_bullet_point(doc, "Google Apps Script Developer Guide (2026). DriveApp & ContentService WebApp Deployment. Retreived from https://developers.google.com/apps-script")
    add_bullet_point(doc, "React Documentation (2026). React 18 Hooks & State Management. Retreived from https://react.dev")
    add_bullet_point(doc, "Tập đoàn TBS Group (2026). Tài liệu Quy trình Vận hành & Tiêu chuẩn Chất lượng Skechers TBS II. Lưu hành nội bộ.")

    # Save modified document
    doc.save(doc_path)
    print(f"Successfully updated document: {doc_path}")

if __name__ == '__main__':
    main()
