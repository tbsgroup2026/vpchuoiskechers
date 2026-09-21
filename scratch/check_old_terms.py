import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx')
full_text = []
for p in doc.paragraphs:
    full_text.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full_text.append(cell.text)
for s in doc.sections:
    for h in [s.header, s.first_page_header, s.even_page_header]:
        if h:
            for p in h.paragraphs:
                full_text.append(p.text)
    for f in [s.footer, s.first_page_footer, s.even_page_footer]:
        if f:
            for p in f.paragraphs:
                full_text.append(p.text)

text = '\n'.join(full_text)

old_terms = [
    'Sóng Thần',
    'Bình Dương',
    'Đồng Nai',
    '1701958307',
    'Bình Đáng',
    'Bình Hòa',
    'Lê Thị Hồng Gấm',
    'Cung Ứng Nhân Lực',
    '0843959131',
    'Vũ Quang Hoan',
    'Đỗ Đức Chí',
    'Văn phòng Chuỗi Skechers'
]

for ot in old_terms:
    cnt = text.count(ot)
    print(f'Old term "{ot}": {cnt}')
