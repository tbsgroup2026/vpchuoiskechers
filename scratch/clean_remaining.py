import docx

def clean_remaining():
    doc_path = r"d:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx"
    doc = docx.Document(doc_path)

    extra_replacements = [
        ("Nhân Kiệt", "TBS Group — VP Chuỗi Skechers"),
        ("nhankiet.vn", "tbsgroup.vn")
    ]

    def process_paragraph(p):
        if not p.text:
            return
        for old_str, new_str in extra_replacements:
            if old_str in p.text:
                for r in p.runs:
                    if old_str in r.text:
                        r.text = r.text.replace(old_str, new_str)
                if old_str in p.text:
                    if len(p.runs) > 0:
                        p.runs[0].text = p.text.replace(old_str, new_str)
                        for r in p.runs[1:]:
                            r.text = ""

    def process_table(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)

    for p in doc.paragraphs:
        process_paragraph(p)

    for t in doc.tables:
        process_table(t)

    for section in doc.sections:
        for p in section.header.paragraphs:
            process_paragraph(p)
        for t in section.header.tables:
            process_table(t)
        for p in section.footer.paragraphs:
            process_paragraph(p)
        for t in section.footer.tables:
            process_table(t)

    doc.save(doc_path)
    print("Cleaned remaining terms.")

if __name__ == "__main__":
    clean_remaining()
