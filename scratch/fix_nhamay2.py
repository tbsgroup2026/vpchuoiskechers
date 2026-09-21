import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def update_to_nhamay2():
    doc = docx.Document(DOC_PATH)
    
    # Define exact target values specified by the user
    TARGET_COMPANY = "Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2"
    TARGET_COMPANY_SHORT = "Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II)"
    TARGET_ADDRESS = "2/434, Khu phố Bình Đáng, Phường Bình Hòa, Thành phố Hồ Chí Minh"
    TARGET_TOPIC = "Xây dựng và phát triển Hệ thống Quản trị số hóa Vận hành Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II Workspace, Gemba Walk 5S Audit & Kaizen IE Engine)"
    
    # -------------------------------------------------------------
    # 1. PARAGRAPH MAPPINGS
    # -------------------------------------------------------------
    replacements = [
        # Topic Title
        (r'XÂY DỰNG VÀ PHÁT TRIỂN HỆ THỐNG QUẢN TRỊ SỐ HÓA.*', TARGET_TOPIC.upper()),
        (r'Xây dựng và phát triển Hệ thống Quản trị số hóa.*', TARGET_TOPIC),
        
        # Company Names
        (r'Công Ty Cổ Phần Thái Bình Kiên Giang \(TBS II - TBS GROUP\)', TARGET_COMPANY_SHORT),
        (r'Công ty Cổ phần Thái Bình Kiên Giang \(TBS II\)', TARGET_COMPANY_SHORT),
        (r'Công Ty Cổ Phần Thái Bình Kiên Giang', TARGET_COMPANY),
        (r'Công ty Cổ phần Thái Bình Kiên Giang', TARGET_COMPANY),
        (r'CÔNG TY CỔ PHẦN THÁI BÌNH KIÊN GIANG \(TBS II\)', TARGET_COMPANY.upper() + " (TBS II)"),
        (r'CÔNG TY CỔ PHẦN THÁI BÌNH KIÊN GIANG', TARGET_COMPANY.upper()),
        (r'Tập đoàn Da Giày TBS Group — Văn phòng Chuỗi Skechers', TARGET_COMPANY_SHORT),
        (r'Công ty TNHH Cung Ứng Nhân Lực TBS Group — VP Chuỗi Skechers', TARGET_COMPANY_SHORT),
        
        # Addresses
        (r'Khu công nghiệp Thạnh Lộc, Xã Thạnh Lộc, Huyện Châu Thành, Tỉnh Kiên Giang, Việt Nam', TARGET_ADDRESS),
        (r'Khu công nghiệp Thạnh Lộc, Xã Thạnh Lộc, Huyện Châu Thành, Tỉnh Kiên Giang', TARGET_ADDRESS),
        (r'KCN Thạnh Lộc, Xã Thạnh Lộc, Huyện Châu Thành, Tỉnh Kiên Giang', TARGET_ADDRESS),
        (r'KCN Thạnh Lộc, Châu Thành, Kiên Giang', TARGET_ADDRESS),
        (r'KCN Thạnh Lộc', 'Khu phố Bình Đáng, Phường Bình Hòa'),
        (r'tỉnh Kiên Giang \(KCN Thạnh Lộc, Châu Thành\)', TARGET_ADDRESS),
        (r'Châu Thành, Kiên Giang', 'Bình Hòa, Thành phố Hồ Chí Minh'),
        (r'Kiên Giang', 'Thành phố Hồ Chí Minh'),

        # Plant details
        (r'Nhà máy KG1, Nhà máy KG2, Nhà máy KG3, Xưởng Đế PU/Cao su', 'Xưởng May, Xưởng Gò Ráp, Xưởng Đế PU/Cao su, Xưởng Đóng gói thuộc Nhà Máy 2 (TBS II)'),
        (r'Tổ hợp Kiên Giang \(TBS II - KG1, KG2, KG3 & Xưởng Đế PU/Cao su\)', 'Cụm Nhà Máy 2 (TBS II)'),
        (r'Tổ hợp Nhà máy Kiên Giang \(TBS II\)', 'Cụm Nhà Máy 2 (TBS II)'),
        (r'Tổ hợp Kiên Giang', 'Nhà Máy 2 (TBS II)'),
        (r'3 nhà máy lớn \(KG1, KG2, KG3\)', 'các xưởng sản xuất quy mô lớn (May, Gò, Ráp, Đế)'),
        (r'KG1', 'Xưởng Gò Ráp'),
        (r'KG2', 'Xưởng May'),
        (r'KG3', 'Xưởng Đóng Gói'),
    ]

    print("Updating paragraphs...")
    for p in doc.paragraphs:
        for old_pattern, new_text in replacements:
            if re.search(old_pattern, p.text):
                p.text = re.sub(old_pattern, new_text, p.text)

    # -------------------------------------------------------------
    # 2. TABLES UPDATE
    # -------------------------------------------------------------
    print("Updating tables...")
    for t_idx, t in enumerate(doc.tables):
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                for old_pattern, new_text in replacements:
                    if re.search(old_pattern, cell.text):
                        cell.text = re.sub(old_pattern, new_text, cell.text)

    # Table 8: Organizational Structure for Nhà Máy 2
    if len(doc.tables) > 8:
        t8 = doc.tables[8]
        if len(t8.rows) >= 4:
            t8.rows[0].cells[0].text = "BAN TỔNG GIÁM ĐỐC CÔNG TY CP ĐẦU TƯ THÁI BÌNH - NHÀ MÁY 2 (TBS II)\n(Điều hành Hoạt động Sản xuất & Chuyển đổi số)"
            for c in t8.rows[0].cells[1:]:
                c.text = t8.rows[0].cells[0].text
            
            t8.rows[1].cells[0].text = "PHÒNG CÔNG NGHỆ THÔNG TIN & CĐS\n(Anh Ngô Hà Thanh An - Quản lý CĐS / HDTT)"
            t8.rows[1].cells[1].text = "PHÒNG KỸ THUẬT IE & QUẢN TRỊ SẢN XUẤT\n(Industrial Engineering & Operations)"
            t8.rows[1].cells[2].text = "PHÒNG QUẢN LÝ CHẤT LƯỢNG & GEMBA 5S\n(QA/QC & Kaizen Engine)"
            t8.rows[1].cells[3].text = "PHÒNG NHÂN SỰ & QUẢN TRỊ NGUỒN LỰC\n(HR & Administration)"
            if len(t8.rows[1].cells) > 4:
                t8.rows[1].cells[4].text = "KHỐI CÁC XƯỞNG SẢN XUẤT NHÀ MÁY 2\n(Xưởng May - Gò - Ráp - Đế PU/Cao su)"

            t8.rows[2].cells[0].text = f"▼ {TARGET_COMPANY_SHORT} – {TARGET_ADDRESS}"
            for c in t8.rows[2].cells[1:]:
                c.text = t8.rows[2].cells[0].text

            t8.rows[3].cells[0].text = "Xưởng Gò Ráp\n(Sản xuất & Gò ráp Giày Skechers)"
            t8.rows[3].cells[1].text = "Xưởng May\n(May Thể thao xuất khẩu)"
            t8.rows[3].cells[2].text = "Xưởng Đóng Gói\n(Đóng gói & Kiểm định Skechers)"
            if len(t8.rows[3].cells) > 3:
                t8.rows[3].cells[3].text = "Xưởng Phụ liệu Đế PU / Cao su"
            if len(t8.rows[3].cells) > 4:
                t8.rows[3].cells[4].text = "Hệ thống Số hóa TBS II Workspace"

    # Table 9: Business Activities Table for Nhà Máy 2
    if len(doc.tables) > 9:
        t9 = doc.tables[9]
        if len(t9.rows) >= 4:
            t9.rows[0].cells[0].text = f"LĨNH VỰC HOẠT ĐỘNG & NĂNG LỰC SẢN XUẤT - {TARGET_COMPANY.upper()}"
            for c in t9.rows[0].cells[1:]:
                c.text = t9.rows[0].cells[0].text

            t9.rows[1].cells[0].text = "Sản xuất Giày Thể thao Skechers xuất khẩu"
            t9.rows[1].cells[1].text = "Sản xuất Phụ liệu & Đế PU / Cao su"
            t9.rows[1].cells[2].text = "Sản xuất Phụ liệu & Đế PU / Cao su"
            t9.rows[1].cells[3].text = "Quản trị Kỹ thuật IE & Định mức SAM"
            if len(t9.rows[1].cells) > 4:
                t9.rows[1].cells[4].text = "Số hóa Vận hành & Gemba Walk 5S Audit"

            t9.rows[2].cells[0].text = "May, gò, ráp và đóng gói giày thể thao Skechers đạt tiêu chuẩn quốc tế"
            t9.rows[2].cells[1].text = "Ép đế PU, lưu hóa đế cao su và cung ứng phụ liệu cho các chuyền sản xuất"
            t9.rows[2].cells[2].text = "Ép đế PU, lưu hóa đế cao su và cung ứng phụ liệu cho các chuyền sản xuất"
            t9.rows[2].cells[3].text = "Tính định mức SAM, tối ưu hóa thao tác chuyền sản xuất và cân bằng chuyền"
            if len(t9.rows[2].cells) > 4:
                t9.rows[2].cells[4].text = "Số hóa thẻ Kanban task, chấm điểm 5S Audit và Kaizen IE trực tuyến"

            t9.rows[3].cells[0].text = f"CƠ CẤU VẬN HÀNH NỘI BỘ - {TARGET_COMPANY_SHORT}"
            for c in t9.rows[3].cells[1:]:
                c.text = t9.rows[3].cells[0].text

            if len(t9.rows) > 4:
                t9.rows[4].cells[0].text = "Khu vực Khối Sản xuất (Xưởng May & Gò Ráp)"
                t9.rows[4].cells[1].text = "Khu vực Khối Sản xuất (Xưởng May & Gò Ráp)"
                t9.rows[4].cells[2].text = "Khu vực Khối Phụ liệu (Xưởng Đế & Đóng gói)"
                t9.rows[4].cells[3].text = "Khu vực Khối Phụ liệu (Xưởng Đế & Đóng gói)"
                if len(t9.rows[4].cells) > 4:
                    t9.rows[4].cells[4].text = "Khối Quản trị Số hóa & IT (TBS II Workspace)"

            if len(t9.rows) > 5:
                t9.rows[5].cells[0].text = f"• {TARGET_COMPANY}\n• Địa chỉ: {TARGET_ADDRESS}\n• Xưởng May – Chuyền may Thể thao xuất khẩu\n• Xưởng Gò Ráp – Chuyền gò ráp Giày Skechers"
                t9.rows[5].cells[1].text = t9.rows[5].cells[0].text
                t9.rows[5].cells[2].text = f"• {TARGET_COMPANY}\n• Địa chỉ: {TARGET_ADDRESS}\n• Xưởng phụ liệu – Ép đế PU & Lưu hóa cao su\n• Xưởng Đóng gói & Kiểm định chất lượng"
                t9.rows[5].cells[3].text = t9.rows[5].cells[2].text
                if len(t9.rows[5].cells) > 4:
                    t9.rows[5].cells[4].text = "• Trung tâm Số hóa Vận hành & CĐS\n• Hệ thống TBS II Workspace & Kaizen Engine\n• Quản lý CĐS: Ngô Hà Thanh An\n• Điện thoại: 0274 3758 888"

    # -------------------------------------------------------------
    # 3. HEADERS & FOOTERS UPDATE
    # -------------------------------------------------------------
    print("Updating headers and footers...")
    for s in doc.sections:
        for h in [s.header, s.first_page_header, s.even_page_header]:
            if h:
                for p in h.paragraphs:
                    for old_pattern, new_text in replacements:
                        if re.search(old_pattern, p.text):
                            p.text = re.sub(old_pattern, new_text, p.text)
        for f in [s.footer, s.first_page_footer, s.even_page_footer]:
            if f:
                for p in f.paragraphs:
                    for old_pattern, new_text in replacements:
                        if re.search(old_pattern, p.text):
                            p.text = re.sub(old_pattern, new_text, p.text)

    # Save document
    doc.save(DOC_PATH)
    print(f"Successfully saved updated document to {DOC_PATH}")

if __name__ == '__main__':
    update_to_nhamay2()
