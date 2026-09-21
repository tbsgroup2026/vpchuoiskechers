import docx
import copy
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def fix_row11_xml():
    doc = docx.Document(DOC_PATH)
    t2 = doc.tables[2]

    # Clone Row 10 XML for Row 11
    tr10 = t2.rows[10]._tr
    tr11 = t2.rows[11]._tr
    new_tr = copy.deepcopy(tr10)
    tr11.getparent().replace(tr11, new_tr)

    # Re-fetch table rows
    t2 = doc.tables[2]
    r11 = t2.rows[11]

    # Set exact text for each cell in row 11
    r11.cells[0].text = "11"
    r11.cells[1].text = "26/10 – 31/10"
    r11.cells[2].text = "26/10 – 31/10"
    r11.cells[3].text = "Tổng kết dự án, bàn giao mã nguồn & hệ thống TBS II Workspace, hoàn thiện và nộp báo cáo thực tập"
    r11.cells[4].text = "Kế hoạch"

    doc.save(DOC_PATH)
    print("Successfully cloned XML and updated Row 11 in Table 2!")

if __name__ == '__main__':
    fix_row11_xml()
