import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def fix_table2():
    doc = docx.Document(DOC_PATH)
    t2 = doc.tables[2]

    # Row 11: Week 11
    r11 = t2.rows[11]
    r11.cells[0].text = "11"
    r11.cells[1].text = "26/10 – 31/10"
    r11.cells[2].text = "26/10 – 31/10"
    r11.cells[3].text = "Tổng kết dự án, bàn giao mã nguồn & hệ thống TBS II Workspace, hoàn thiện và nộp báo cáo thực tập"
    r11.cells[4].text = "Kế hoạch"

    # Row 12: Signatures
    r12 = t2.rows[12]
    sig_text = "Thành phố Hồ Chí Minh, ngày 12 tháng 09 năm 2026\n\nSinh viên (Ký và ghi rõ họ tên)                Xác nhận của đơn vị thực tập"
    r12.cells[0].text = "Cán bộ hướng dẫn\n(Ký và ghi rõ họ tên)"
    r12.cells[1].text = "Cán bộ hướng dẫn\n(Ký và ghi rõ họ tên)"
    r12.cells[2].text = sig_text
    r12.cells[3].text = sig_text
    r12.cells[4].text = sig_text

    # Also clean up any accidental "111 tuần" regex artifacts in paragraphs
    for p in doc.paragraphs:
        if '111 tuần' in p.text:
            p.text = p.text.replace('111 tuần (17/08 – 31/10/2026)', '12')
            p.text = p.text.replace('111 tuần', '11 tuần')
        if 'tháng 11 tuần' in p.text:
            p.text = p.text.replace('tháng 11 tuần (17/08 – 31/10/2026)', 'tháng 09')

    doc.save(DOC_PATH)
    print("Table 2 updated cleanly!")

if __name__ == '__main__':
    fix_table2()
