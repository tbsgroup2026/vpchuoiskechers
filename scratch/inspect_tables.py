import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx')
print(f"Total tables: {len(doc.tables)}")

for i, t in enumerate(doc.tables):
    print(f"\n=== Table {i} ({len(t.rows)} rows, {len(t.columns)} cols) ===")
    for r_idx in range(min(4, len(t.rows))):
        row = t.rows[r_idx]
        cells_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"  Row {r_idx}: {' | '.join(cells_text)[:140]}")
