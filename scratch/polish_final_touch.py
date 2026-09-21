import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def final_touch():
    doc = docx.Document(DOC_PATH)
    
    final_transforms = [
        (r'Dịch vụ cung ứng lao động \(Outsourcing Services\):.*', 
         'Khối Sản xuất Da Giày Skechers xuất khẩu: May, gò, ráp và hoàn thiện các dòng sản phẩm giày thể thao đạt tiêu chuẩn xuất khẩu sang thị trường Mỹ và Châu Âu.'),
        (r'Về tuyển dụng và cung ứng lao động:.*', 
         'Về Vận hành & Quản lý Sản xuất: Đảm bảo tiến độ chuyền may/gò/ráp, tối ưu thời gian định mức SAM, thực hiện kiểm định 5S Gemba Walk và thúc đẩy các phong trào cải tiến Kaizen IE.'),
        (r'người tuyển dụng', 'Trưởng phòng / Kỹ sư IE'),
        (r'hoàn tất tuyển dụng', 'hoàn tất quy trình nghiệm thu & duyệt Kaizen IE'),
        (r'phê duyệt tuyển dụng', 'phê duyệt Kaizen IE & nghiệm thu task'),
        (r'tin tuyển dụng', 'thẻ nhiệm vụ Kanban'),
        (r'job matching', 'Kaizen IE matching'),
    ]

    print("Applying final touches to paragraphs...")
    for p in doc.paragraphs:
        for old_patt, new_t in final_transforms:
            if re.search(old_patt, p.text):
                p.text = re.sub(old_patt, new_t, p.text)

    print("Applying final touches to tables...")
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for old_patt, new_t in final_transforms:
                    if re.search(old_patt, c.text):
                        c.text = re.sub(old_patt, new_t, c.text)

    doc.save(DOC_PATH)
    print("Final touches applied successfully!")

if __name__ == '__main__':
    final_touch()
