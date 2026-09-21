import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def remove_de_mui_go():
    doc = docx.Document(DOC_PATH)
    
    replacements = [
        (r'Xưởng đế-mũi-gò', 'Xưởng May, Xưởng Gò Ráp và Xưởng Đế PU/Cao su'),
        (r'xưởng đế-mũi-gò', 'xưởng May, xưởng Gò Ráp và xưởng Đế PU/Cao su'),
        (r'Đế-mũi-gò', 'Gò Ráp và Đế PU/Cao su'),
        (r'đế-mũi-gò', 'gò ráp và đế PU/cao su'),
        (r'Đế - Mũi - Gò', 'Gò Ráp & Đế PU/Cao su'),
        (r'đế - mũi - gò', 'gò ráp & đế PU/cao su'),
        (r'Đế-mũi', 'Đế PU/Cao su'),
        (r'đế-mũi', 'đế PU/cao su'),
    ]

    print("Checking and replacing any 'Xưởng đế-mũi-gò' references...")
    modified_paragraphs = 0
    modified_table_cells = 0

    for p in doc.paragraphs:
        for old_patt, new_t in replacements:
            if re.search(old_patt, p.text, re.IGNORECASE):
                p.text = re.sub(old_patt, new_t, p.text, flags=re.IGNORECASE)
                modified_paragraphs += 1

    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for old_patt, new_t in replacements:
                    if re.search(old_patt, c.text, re.IGNORECASE):
                        c.text = re.sub(old_patt, new_t, c.text, flags=re.IGNORECASE)
                        modified_table_cells += 1

    for s in doc.sections:
        for h in [s.header, s.first_page_header, s.even_page_header]:
            if h:
                for p in h.paragraphs:
                    for old_patt, new_t in replacements:
                        if re.search(old_patt, p.text, re.IGNORECASE):
                            p.text = re.sub(old_patt, new_t, p.text, flags=re.IGNORECASE)
        for f in [s.footer, s.first_page_footer, s.even_page_footer]:
            if f:
                for p in f.paragraphs:
                    for old_patt, new_t in replacements:
                        if re.search(old_patt, p.text, re.IGNORECASE):
                            p.text = re.sub(old_patt, new_t, p.text, flags=re.IGNORECASE)

    doc.save(DOC_PATH)
    print(f"Replacement complete! Modified {modified_paragraphs} paragraphs and {modified_table_cells} table cells.")

if __name__ == '__main__':
    remove_de_mui_go()
