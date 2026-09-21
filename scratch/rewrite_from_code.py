import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def rewrite_based_on_codebase():
    doc = docx.Document(DOC_PATH)
    
    # -------------------------------------------------------------
    # EXACT CODEBASE REPLACEMENT MAP (STRICTLY FROM d:\Work\TBS II)
    # -------------------------------------------------------------
    code_replacements = [
        # System & Topic Name
        (r'XÂY DỰNG VÀ PHÁT TRIỂN HỆ THỐNG QUẢN TRỊ SỐ HÓA.*',
         'XÂY DỰNG VÀ PHÁT TRIỂN HỆ THỐNG QUẢN TRỊ SỐ HÓA VẬN HÀNH CÔNG TY CP ĐẦU TƯ THÁI BÌNH - NHÀ MÁY 2 (TBS II WORKSPACE, GEMBA WALK 5S AUDIT & KAIZEN IE ENGINE)'),
        (r'Xây dựng và phát triển Hệ thống Quản trị số hóa.*',
         'Xây dựng và phát triển Hệ thống Quản trị số hóa Vận hành Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II Workspace, Gemba Walk 5S Audit & Kaizen IE Engine)'),
        (r'TBS Work Hub & Gemba Kaizen', 'TBS II Workspace (Gemba Walk 5S Audit & Kaizen IE Engine)'),
        (r'TBS Work Hub', 'TBS II Workspace'),

        # Company, Location, Institutional Info
        (r'Công Ty Cổ Phần Thái Bình Kiên Giang.*', 'Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II)'),
        (r'Công ty Cổ phần Thái Bình Kiên Giang.*', 'Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II)'),
        (r'Tập đoàn Da Giày TBS Group — Văn phòng Chuỗi Skechers', 'Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II)'),
        (r'Công ty TNHH Cung Ứng Nhân Lực TBS Group — VP Chuỗi Skechers', 'Công Ty CP Đầu Tư Thái Bình - Nhà Máy 2 (TBS II)'),
        (r'Ngô Hà Thanh An', 'Dư Thị Thanh Tình'),
        (r'PGĐ\. Ngô Hà Thanh An', 'Dư Thị Thanh Tình (TP. Hành chính & Nhân sự)'),

        # Operational Workflows derived 100% from codebase API routes & SQL migrations
        (r'luồng đăng ký/đăng nhập đa kênh.*', 
         'quy trình đăng nhập xác thực MSNV, phân quyền RBAC/ABAC (`0001_concurrency_rbac.sql`), khởi tạo thẻ nhiệm vụ Kanban (`/api/task-cards`), quy trình nộp đề xuất Kaizen 5 bước (`/api/ci-kaizen`) và kiểm định 5S Gemba Walk Audit (`/api/maintenance`).'),
        (r'sơ đồ phễu tuyển dụng.*', 
         'quy trình 5 bước nộp & duyệt Kaizen IE (`ci-kaizen`): (1) CBCNV nộp Kaizen 5 bước -> (2) AI so sánh trùng lặp (`/api/ai/compare-kaizen`) -> (3) Trưởng phòng phân công Reviewer (`/api/ci-kaizen/assign-reviewer`) -> (4) Chuyên gia đánh giá (`/api/ci-kaizen/expert-evaluations`) -> (5) Kỹ sư IE xác nhận định mức (`/api/ci-kaizen/[id]/ie-confirm`) và xếp hạng (`/api/ci-kaizen/ranking`).'),
        (r'Retention Features và tuyển dụng thông minh', 
         'Tối ưu hóa UI/UX PWA, phân công Reviewer & duyệt Kaizen đa cấp, tích hợp kiểm định thiết bị MMTB (`0006_mmtb_integration.sql`) & thông báo đa kênh (`/api/push/subscribe`).'),
        (r'quản lý tin đăng tuyển dụng', 'quản lý thẻ nhiệm vụ Kanban (`/api/task-boards`) & theo dõi bảng xếp hạng Kaizen (`/api/ci-kaizen/ranking`)'),
        (r'nghiệm thu & duyệt kết quả công việc và tuyển dụng', 'kiểm định 5S Gemba Walk Audit & nghiệm thu đề xuất Kaizen IE (`ci-kaizen`)'),
        (r'nghiệm thu & duyệt kết quả công việc', 'kiểm định 5S Gemba Walk & duyệt Kaizen IE'),
    ]

    print("Replacing document text strictly according to codebase...")
    for p in doc.paragraphs:
        for old_patt, new_t in code_replacements:
            if re.search(old_patt, p.text):
                p.text = re.sub(old_patt, new_t, p.text)

    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for old_patt, new_t in code_replacements:
                    if re.search(old_patt, c.text):
                        c.text = re.sub(old_patt, new_t, c.text)

    for s in doc.sections:
        for h in [s.header, s.first_page_header, s.even_page_header]:
            if h:
                for p in h.paragraphs:
                    for old_patt, new_t in code_replacements:
                        if re.search(old_patt, p.text):
                            p.text = re.sub(old_patt, new_t, p.text)
        for f in [s.footer, s.first_page_footer, s.even_page_footer]:
            if f:
                for p in f.paragraphs:
                    for old_patt, new_t in code_replacements:
                        if re.search(old_patt, p.text):
                            p.text = re.sub(old_patt, new_t, p.text)

    doc.save(DOC_PATH)
    print(f"Successfully saved updated document to {DOC_PATH}")

if __name__ == '__main__':
    rewrite_based_on_codebase()
