import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def update_advisor():
    doc = docx.Document(DOC_PATH)
    
    OLD_ADVISOR = "Ngô Hà Thanh An"
    NEW_ADVISOR_FULL = "Dư Thị Thanh Tình (TP. Hành chính & Nhân sự)"
    NEW_ADVISOR_NAME = "Dư Thị Thanh Tình"
    
    replacements = [
        (r'Người hướng dẫn tại đơn vị: PGĐ\. Ngô Hà Thanh An', f'Người hướng dẫn tại đơn vị: {NEW_ADVISOR_FULL}'),
        (r'Người hướng dẫn tại đơn vị: Ngô Hà Thanh An', f'Người hướng dẫn tại đơn vị: {NEW_ADVISOR_FULL}'),
        (r'Cán bộ hướng dẫn: Ngô Hà Thanh An', f'Cán bộ hướng dẫn: {NEW_ADVISOR_FULL}'),
        (r'Quản lý CĐS: Ngô Hà Thanh An', f'TP. Hành chính & Nhân sự: {NEW_ADVISOR_NAME}'),
        (r'Anh Ngô Hà Thanh An - Quản lý CĐS / HDTT', f'Bà {NEW_ADVISOR_FULL} - Cán bộ hướng dẫn'),
        (r'Ngô Hà Thanh An', NEW_ADVISOR_NAME),
    ]

    print("Updating paragraphs...")
    for p in doc.paragraphs:
        for old_pattern, new_text in replacements:
            if re.search(old_pattern, p.text):
                p.text = re.sub(old_pattern, new_text, p.text)

    print("Updating tables...")
    for t_idx, t in enumerate(doc.tables):
        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                for old_pattern, new_text in replacements:
                    if re.search(old_pattern, cell.text):
                        cell.text = re.sub(old_pattern, new_text, cell.text)

    # Table 8 Row 1 Cell 0
    if len(doc.tables) > 8:
        t8 = doc.tables[8]
        if len(t8.rows) > 1:
            t8.rows[1].cells[0].text = f"PHÒNG HÀNH CHÍNH & NHÂN SỰ\n(Bà {NEW_ADVISOR_FULL} - Cán bộ hướng dẫn)"

    # Table 9 Row 5 Cells
    if len(doc.tables) > 9:
        t9 = doc.tables[9]
        if len(t9.rows) > 5:
            for c_idx in range(4, len(t9.rows[5].cells)):
                t9.rows[5].cells[c_idx].text = f"• Trung tâm Số hóa Vận hành & CĐS\n• Hệ thống TBS II Workspace & Kaizen Engine\n• TP. Hành chính & Nhân sự: {NEW_ADVISOR_NAME}\n• Điện thoại: 0274 3758 888"

    print("Updating headers and footers...")
    for s in doc.sections:
        for h in [s.header, s.first_page_header, s.even_page_header]:
            if h:
                for p in h.paragraphs:
                    for old_pattern, new_text in replacements:
                        if re.search(old_pattern, p.text):
                            p.text = re.sub(old_pattern, new_text, p.text)
        for f in [s.footer, s.first_page_footer, s.even_page_footer]:
            if f:
                for p in f.paragraphs:
                    for old_pattern, new_text in replacements:
                        if re.search(old_pattern, p.text):
                            p.text = re.sub(old_pattern, new_text, p.text)

    doc.save(DOC_PATH)
    print(f"Successfully saved updated document to {DOC_PATH}")

if __name__ == '__main__':
    update_advisor()
