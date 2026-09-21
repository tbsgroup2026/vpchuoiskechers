import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx'

def update_journal():
    doc = docx.Document(DOC_PATH)
    
    # -------------------------------------------------------------
    # 1. UPDATE PARAGRAPHS WITH NEW DATES (17/08/2026 - 31/10/2026)
    # -------------------------------------------------------------
    date_replacements = [
        (r'18 / 05 / 2026 đến 26 / 07 / 2026', '17 / 08 / 2026 đến 31 / 10 / 2026'),
        (r'18/5/2026 đến ngày 26/7/2026', '17/8/2026 đến ngày 31/10/2026'),
        (r'18/5 đến 26/7/2026', '17/8 đến 31/10/2026'),
        (r'18/5 đến 26/7', '17/8 đến 31/10'),
        (r'18/05 – 22/05/2026', '17/08 – 23/08/2026'),
        (r'25/05 – 29/05/2026', '24/08 – 30/08/2026'),
        (r'01/06 – 05/06/2026', '31/08 – 06/09/2026'),
        (r'01/06 – 07/06/2026', '31/08 – 06/09/2026'),
        (r'08/06 – 12/06/2026', '07/09 – 13/09/2026'),
        (r'08/06 – 14/06/2026', '07/09 – 13/09/2026'),
        (r'15/06 – 26/06/2026', '14/09 – 27/09/2026'),
        (r'29/06 – 10/07/2026', '28/09 – 11/10/2026'),
        (r'13/07 – 26/07/2026', '12/10 – 31/10/2026'),
        (r'10 tuần', '11 tuần'),
        (r'2 tháng', '11 tuần (17/08 – 31/10/2026)'),
        (r'ngày 27  tháng 7  năm 2026', 'ngày 12 tháng 09 năm 2026'),
        (r'ngày 27 tháng 7 năm 2026', 'ngày 12 tháng 09 năm 2026'),
        (r'ngày….tháng….năm……', 'ngày 12 tháng 09 năm 2026'),
    ]

    print("Updating paragraph dates...")
    for p in doc.paragraphs:
        for old_patt, new_t in date_replacements:
            if re.search(old_patt, p.text):
                p.text = re.sub(old_patt, new_t, p.text)

    # -------------------------------------------------------------
    # 2. UPDATE TABLE 2 (NHẬT KÝ THỰC TẬP TABLE)
    # -------------------------------------------------------------
    print("Updating Table 2 (Nhật ký thực tập)...")
    if len(doc.tables) > 2:
        t2 = doc.tables[2]
        
        # Header Row (Row 0)
        row0 = t2.rows[0]
        row0.cells[0].text = "Tuần lễ"
        row0.cells[1].text = "Từ ngày 17/8 đến 31/10"
        row0.cells[2].text = "Từ ngày 17/8 đến 31/10"
        row0.cells[3].text = "Nội dung"
        row0.cells[4].text = "Ghi chú"

        # Weekly Rows Data
        weeks_data = [
            ("1", "17/8 – 23/8", "17/8 – 23/8", "Tìm hiểu tổng quan công ty, khảo sát thực tế quy trình Gemba Walk 5S tại Nhà Máy 2, phân tích yêu cầu & lập kế hoạch thực tập", "Hoàn thành"),
            ("2", "24/8 – 30/8", "24/8 – 30/8", "Xây dựng các chức năng cốt lõi (Đăng ký/đăng nhập đa kênh, thẻ Kanban công việc, đồng bộ API Surface kép trên Express và Workers)", "Hoàn thành"),
            ("3", "31/8 – 06/9", "31/8 – 06/9", "Phát triển AI Automation, AI IE Chatbot, AI CV Analysis, Dashboard Thống kê Admin & cơ chế bảo mật Rate Limiting", "Hoàn thành"),
            ("4", "07/9 – 13/9", "07/9 – 13/9", "Triển khai Gemba Walk 5S Audit & Kaizen IE Engine, phân công Reviewer & duyệt Kaizen đa cấp, tích hợp PayOS/Zalo, báo cáo tiến độ đợt 1", "Đang thực hiện (Tính đến 12/9)"),
            ("5", "14/9 – 20/9", "14/9 – 20/9", "Tái cấu trúc giao diện PWA, chuẩn hóa Design System (#0A8043) & mở rộng tính năng quản lý dành cho Trưởng phòng/Kỹ sư IE", "Kế hoạch"),
            ("6", "21/9 – 27/9", "21/9 – 27/9", "Phát triển các tính năng Retention, Video Recruitment, Career DNA & thuật toán đề xuất việc làm bằng Cosine Similarity", "Kế hoạch"),
            ("7", "28/9 – 04/10", "28/9 – 04/10", "Cấu hình hạ tầng Edge Computing (Cloudflare D1/R2), bảo mật hệ thống (Security Headers, PBKDF2) & triển khai CI/CD Actions", "Kế hoạch"),
            ("8", "05/10 – 11/10", "05/10 – 11/10", "Tối ưu hiệu năng toàn diện (API latency < 50ms), kiểm thử tải (Stress testing) & khắc phục lỗi sản phẩm trên production", "Kế hoạch"),
            ("9", "12/10 – 18/10", "12/10 – 18/10", "Viết tài liệu hướng dẫn sử dụng hệ thống số hóa TBS II Workspace & tài liệu quy trình Gemba Walk 5S cho CBCNV", "Kế hoạch"),
            ("10", "19/10 – 25/10", "19/10 – 25/10", "Tập huấn ứng dụng số hóa cho cán bộ quản lý & công nhân tại Nhà Máy 2, kiểm thử nghiệm thu người dùng (UAT)", "Kế hoạch"),
            ("11", "26/10 – 31/10", "26/10 – 31/10", "Tổng kết dự án, bàn giao mã nguồn & hệ thống TBS II Workspace, hoàn thiện và nộp báo cáo thực tập doanh nghiệp", "Kế hoạch")
        ]

        # Check existing rows in t2
        # Current rows: Row 0 (header), Rows 1-10 (Tuần 1-10), Row 11 (signature)
        # We need 11 weekly rows + header + signature = 13 rows total.
        # Let's see if we need to insert a row before signature row (Row 11).
        
        # Update existing rows 1 to 10
        for idx in range(10):
            w_num, w_date1, w_date2, w_content, w_note = weeks_data[idx]
            r = t2.rows[idx + 1]
            r.cells[0].text = w_num
            r.cells[1].text = w_date1
            r.cells[2].text = w_date2
            r.cells[3].text = w_content
            r.cells[4].text = w_note

        # If t2 has 12 rows, t2.rows[11] is signature row. Let's add row for week 11 before signature.
        if len(t2.rows) == 12:
            # Save signature content
            sig_cells = [c.text for c in t2.rows[11].cells]
            
            # Use row 11 for week 11
            w_num, w_date1, w_date2, w_content, w_note = weeks_data[10]
            r11 = t2.rows[11]
            r11.cells[0].text = w_num
            r11.cells[1].text = w_date1
            r11.cells[2].text = w_date2
            r11.cells[3].text = w_content
            r11.cells[4].text = w_note
            
            # Add new row 12 for signature
            new_r = t2.add_row()
            for c_idx in range(len(new_r.cells)):
                if c_idx < len(sig_cells):
                    new_r.cells[c_idx].text = sig_cells[c_idx]

        # Update date text in signature row
        sig_row = t2.rows[-1]
        for c in sig_row.cells:
            if 'tháng' in c.text:
                c.text = re.sub(r'ngày\s*\d+\s*tháng\s*\d+\s*năm\s*\d+', 'ngày 12 tháng 09 năm 2026', c.text)
                c.text = re.sub(r'ngày….tháng….năm……', 'ngày 12 tháng 09 năm 2026', c.text)

    # -------------------------------------------------------------
    # 3. UPDATE OTHER TABLES DATES AND FOOTERS/HEADERS
    # -------------------------------------------------------------
    print("Updating other tables and headers/footers...")
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for old_patt, new_t in date_replacements:
                    if re.search(old_patt, c.text):
                        c.text = re.sub(old_patt, new_t, c.text)

    for s in doc.sections:
        for h in [s.header, s.first_page_header, s.even_page_header]:
            if h:
                for p in h.paragraphs:
                    for old_patt, new_t in date_replacements:
                        if re.search(old_patt, p.text):
                            p.text = re.sub(old_patt, new_t, p.text)
        for f in [s.footer, s.first_page_footer, s.even_page_footer]:
            if f:
                for p in f.paragraphs:
                    for old_patt, new_t in date_replacements:
                        if re.search(old_patt, p.text):
                            p.text = re.sub(old_patt, new_t, p.text)

    doc.save(DOC_PATH)
    print(f"Successfully saved updated document to {DOC_PATH}")

if __name__ == '__main__':
    update_journal()
