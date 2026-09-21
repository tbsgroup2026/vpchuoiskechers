import docx
import os
import re

def update_kiengiang_doc():
    doc_path = r"d:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx"
    print("Loading docx file for TBS II - Kiên Giang update:", doc_path)
    doc = docx.Document(doc_path)

    # Ordered mapping from general/incorrect terms to exact Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II) terms
    replacements = [
        # Topic & Cover Titles
        ("XÂY DỰNG VÀ PHÁT TRIỂN HỆ THỐNG QUẢN TRỊ SỐ HÓA VĂN PHÒNG CHUỖI SKECHERS (HỆ THỐNG TBS WORKSPACE, MODULE GEMBA WALK 5S VÀ ĐỘNG CƠ KAIZEN IE)",
         "XÂY DỰNG VÀ PHÁT TRIỂN HỆ THỐNG QUẢN TRỊ SỐ HÓA VẬN HÀNH CÔNG TY CỔ PHẦN THÁI BÌNH KIÊN GIANG (HỆ THỐNG TBS II WORKSPACE, MODULE GEMBA WALK 5S VÀ ĐỘNG CƠ KAIZEN IE)"),
        ("Xây dựng và phát triển Hệ thống Quản trị số hóa Văn phòng Chuỗi Skechers (TBS Work Dashboard, Gemba Walk 5S Audit & Kaizen IE Engine)",
         "Xây dựng và phát triển Hệ thống Quản trị số hóa Vận hành Công ty Cổ phần Thái Bình Kiên Giang (TBS II Workspace, Gemba Walk 5S Audit & Kaizen IE Engine)"),
        ("Xây dựng và phát triển Hệ thống Quản trị số hóa Văn phòng Chuỗi Skechers (TBS Work Hub)",
         "Xây dựng và phát triển Hệ thống Quản trị số hóa Vận hành Công ty Cổ phần Thái Bình Kiên Giang (TBS II Workspace)"),

        # Company Names
        ("TẬP ĐOÀN DA GIÀY TBS GROUP — VĂN PHÒNG CHUỖI SKECHERS", "CÔNG TY CỔ PHẦN THÁI BÌNH KIÊN GIANG (TBS GROUP - TỔ HỢP KIÊN GIANG / TBS II)"),
        ("Tập đoàn Da Giày TBS Group — Văn phòng Chuỗi Skechers", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS Group - Tổ Hợp Kiên Giang / TBS II)"),
        ("Tập đoàn Da Giày TBS Group (VP Chuỗi Skechers)", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS Group - TBS II)"),
        ("Tập đoàn Da Giày TBS Group", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS Group)"),
        ("Văn phòng Chuỗi Skechers (VP Chuỗi)", "Khối Quản Trị Vận Hành Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II)"),
        ("Văn phòng Chuỗi Skechers", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II)"),
        ("Văn Phòng Chuỗi Skechers", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II)"),
        ("VP Chuỗi Skechers", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II)"),
        ("VP Chuỗi", "Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II)"),

        # Addresses & Contact Info
        ("KCN Sóng Thần 1, Dĩ An, Bình Dương / VP Chuỗi TP.HCM", "Khu công nghiệp Thạnh Lộc, Xã Thạnh Lộc, Huyện Châu Thành, Tỉnh Kiên Giang"),
        ("KCN Sóng Thần 1, Dĩ An, Bình Dương", "Khu công nghiệp Thạnh Lộc, Xã Thạnh Lộc, Huyện Châu Thành, Tỉnh Kiên Giang"),
        ("Đường số 1, Khu công nghiệp Sóng Thần 1, Dĩ An, Bình Dương / VP Chuỗi TP.HCM", "Khu công nghiệp Thạnh Lộc, Xã Thạnh Lộc, Huyện Châu Thành, Tỉnh Kiên Giang"),
        ("3700123456", "1701588998"),
        ("(0274) 3790 388", "(0297) 3922 888"),
        ("(0274) 3790 389", "(0297) 3922 889"),

        # General TBS text adjustments
        ("TBS Work Hub", "TBS II Workspace"),
        ("TBS Work Dashboard", "TBS II Dashboard"),
        ("TBS Work", "TBS II Work")
    ]

    def apply_rep(text):
        res = text
        for old_str, new_str in replacements:
            res = res.replace(old_str, new_str)
        return res

    def process_p(p):
        if not p.text or not p.text.strip():
            return
        
        # Run-level replace
        for r in p.runs:
            if r.text:
                replaced = apply_rep(r.text)
                if replaced != r.text:
                    r.text = replaced

        # Paragraph-level fallback if text split across runs
        full_text = p.text
        replaced_full = apply_rep(full_text)
        if replaced_full != full_text:
            if len(p.runs) > 0:
                p.runs[0].text = replaced_full
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = replaced_full

    def process_t(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_p(p)

    print(f"Updating {len(doc.paragraphs)} main paragraphs...")
    for p in doc.paragraphs:
        process_p(p)

    print(f"Updating {len(doc.tables)} tables...")
    for t in doc.tables:
        process_t(t)

    print(f"Updating {len(doc.sections)} sections (headers & footers)...")
    for section in doc.sections:
        for p in section.header.paragraphs:
            process_p(p)
        for t in section.header.tables:
            process_t(t)
        for p in section.footer.paragraphs:
            process_p(p)
        for t in section.footer.tables:
            process_t(t)

    doc.save(doc_path)
    print("Successfully updated docx file with Công Ty Cổ Phần Thái Bình Kiên Giang (TBS II) details!")

if __name__ == "__main__":
    update_kiengiang_doc()
