import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'd:\Work\TBS II\222480201738 _PhamNguyenAnhHuy (1).docx')
full_text = []
for p in doc.paragraphs:
    full_text.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full_text.append(cell.text)

text = '\n'.join(full_text)

terms = ['đế-mũi-gò', 'đế mũi gò', 'mũi gò', 'mũi-gò', 'Xưởng Đế PU', 'Xưởng Gò Ráp', 'Xưởng May', 'Xưởng Đóng Gói']
for term in terms:
    print(f'Term "{term}": {text.count(term)}')
