import docx
import sys
import os
import zipfile
import xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = '2224802010738_PhamNguyenAnhHuy.docx'

def scan_all_xml():
    print(f"=== DEEP SCANNING ALL XML IN {DOCX_PATH} ===")
    forbidden = ['việc làm xanh', 'viec lam xanh', 'nhân kiệt', 'nhan kiet', 'cung ứng nhân lực', 'cung ung nhan luc']
    
    with zipfile.ZipFile(DOCX_PATH, 'r') as z:
        for filename in z.namelist():
            if filename.endswith(('.xml', '.rels')):
                content = z.read(filename).decode('utf-8', errors='ignore')
                content_lower = content.lower()
                for f in forbidden:
                    if f in content_lower:
                        # find occurrence count
                        cnt = content_lower.count(f)
                        print(f"[FORBIDDEN] In {filename}: '{f}' found {cnt} times!")

def inspect_structure():
    doc = docx.Document(DOCX_PATH)
    print(f"\nDocument Paragraphs: {len(doc.paragraphs)}")
    print(f"Document Tables: {len(doc.tables)}")
    
    print("\n--- TABLE 2 (Internship Plan) ---")
    t2 = doc.tables[2]
    for r_idx, row in enumerate(t2.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {cells}")

    print("\n--- ALL PARAGRAPHS & STYLES ---")
    with open("docx_paragraphs_dump.txt", "w", encoding="utf-8") as out:
        for idx, p in enumerate(doc.paragraphs):
            out.write(f"[{idx:03d}] ({p.style.name}) {p.text}\n")
    print("Dumped all paragraphs to docx_paragraphs_dump.txt")

if __name__ == '__main__':
    scan_all_xml()
    inspect_structure()
