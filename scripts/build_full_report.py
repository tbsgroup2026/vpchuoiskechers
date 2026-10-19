# -*- coding: utf-8 -*-
"""
scripts/build_full_report.py
Full automated report generator for 2224802010738_PhamNguyenAnhHuy.docx
"""

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import sys
import os
import zipfile

sys.stdout.reconfigure(encoding='utf-8')

DOCX_IN = '2224802010738_PhamNguyenAnhHuy_backup.docx'
DOCX_OUT = '2224802010738_PhamNguyenAnhHuy.docx'

print("=== STARTING FULL REPORT BUILD ===")
doc = Document(DOCX_IN)

# -------------------------------------------------------------
# Formatting helpers
# -------------------------------------------------------------
def format_p(p, text, style='Normal', space_before=3, space_after=6, line_spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, font_size=13):
    p.text = text
    p.style = style
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.italic = italic

def set_cell_props(cell, text, bold=False, font_size=10, align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if fill_hex:
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

def add_caption(p, caption_text):
    p.text = caption_text
    p.style = 'Caption'
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.italic = True
        r.font.bold = True

def add_highlighted_run(p, text, is_highlight=False):
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    if is_highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        run.font.bold = True

# -------------------------------------------------------------
# 1. TABLE 2 (10-WEEK PLAN)
# -------------------------------------------------------------
print("Step 1: Populating Table 2 (10-Week Plan)...")
t2 = doc.tables[2]

weeks_plan = [
    ("1", "03/8 – 08/8", "Tìm hiểu và phân tích tổng quan dự án"),
    ("2", "10/8 – 15/8", "Xây dựng và phát triển các chức năng hệ thống"),
    ("3", "17/8 – 22/8", "Xây dựng phân hệ Quản lý công việc cá nhân (My Tasks), Bảng công việc phòng ban (Kanban) và Dự án liên phòng ban (Projects)"),
    ("4", "24/8 – 29/8", "Xây dựng phân hệ Kiểm soát hiện trường GEMBA Walk và Đăng ký sáng kiến cải tiến liên tục (CI-Kaizen)"),
    ("5", "31/8 – 05/9", "Xây dựng phân hệ Quản lý sản lượng PPH nhà máy và tích hợp dữ liệu Máy móc thiết bị (MMTB)"),
    ("6", "07/9 – 12/9", "Tích hợp hệ thống thông báo đa kênh: Web Push (VAPID), Zalo Notification Service và Service Worker nền tảng"),
    ("7", "14/9 – 19/9", "Hoàn thiện hệ thống phân quyền đa cấp (RBAC), bảo mật xác thực (JWT/Session) và cơ chế Audit Log / Idempotency"),
    ("8", "21/9 – 26/9", "Xây dựng cơ chế sao lưu tự động (Incremental Backup sang Google Drive), quản lý file media Cloudflare R2 và tối ưu nén ảnh client"),
    ("9", "28/9 – 03/10", "Kiểm thử toàn diện hệ thống: UAT 100 người dùng, kiểm thử bảo mật đa vai trò, kiểm thử API và sửa lỗi hồi quy"),
    ("10", "05/10 – 10/10", "Triển khai hệ thống lên Cloudflare Workers (Production), đánh giá hiệu năng và hoàn thiện báo cáo thực tập"),
]

for idx, (wn, wd, wc) in enumerate(weeks_plan, start=1):
    row = t2.rows[idx]
    set_cell_props(row.cells[0], wn, bold=True, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_props(row.cells[1], wd, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    if len(row.cells) > 2 and row.cells[2] != row.cells[1]:
        set_cell_props(row.cells[2], wd, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    content_cell = row.cells[3] if len(row.cells) > 3 else row.cells[-2]
    set_cell_props(content_cell, wc, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    note_cell = row.cells[4] if len(row.cells) > 4 else row.cells[-1]
    set_cell_props(note_cell, "", bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

# -------------------------------------------------------------
# 2. TABLE 10 (ABBREVIATIONS - 31 TERMS)
# -------------------------------------------------------------
print("Step 2: Populating Table 10 (31 Abbreviations)...")
t10 = doc.tables[10]

abbr_list = [
    ("1", "API", "Giao diện lập trình ứng dụng (Application Programming Interface)"),
    ("2", "AQL", "Giới hạn chất lượng chấp nhận (Acceptance Quality Limit)"),
    ("3", "BOLA", "Lỗ hổng kiểm soát truy cập cấp đối tượng (Broken Object Level Authorization)"),
    ("4", "CBCNV", "Cán bộ công nhân viên"),
    ("5", "CI", "Cải tiến liên tục (Continuous Improvement)"),
    ("6", "CI/CD", "Tích hợp liên tục và triển khai liên tục (Continuous Integration / Continuous Deployment)"),
    ("7", "CSP", "Chính sách bảo mật nội dung (Content Security Policy)"),
    ("8", "CSS", "Ngôn ngữ định kiểu tầng (Cascading Style Sheets)"),
    ("9", "D1", "Cơ sở dữ liệu quan hệ SQL phân tán tại mạng biên của Cloudflare"),
    ("10", "DOM", "Mô hình đối tượng tài liệu (Document Object Model)"),
    ("11", "ERD", "Sơ đồ quan hệ thực thể (Entity Relationship Diagram)"),
    ("12", "F5", "Thao tác làm mới tải lại trang trên trình duyệt web"),
    ("13", "GEMBA", "Hiện trường sản xuất thực tế theo triết lý Lean Manufacturing"),
    ("14", "HĐQT", "Hội đồng quản trị"),
    ("15", "HTML", "Ngôn ngữ đánh dấu siêu văn bản (HyperText Markup Language)"),
    ("16", "HTTP", "Giao thức truyền tải siêu văn bản (HyperText Transfer Protocol)"),
    ("17", "IDOR", "Tham chiếu đối tượng trực tiếp không an toàn (Insecure Direct Object References)"),
    ("18", "IE", "Kỹ thuật công nghiệp và định mức lao động (Industrial Engineering)"),
    ("19", "JWT", "Chuỗi mã định danh phiên an toàn (JSON Web Token)"),
    ("20", "KPI", "Chỉ số đánh giá hiệu quả công việc trọng yếu (Key Performance Indicator)"),
    ("21", "MMTB", "Máy móc thiết bị"),
    ("22", "MTTD", "Thời gian trung bình phát hiện sự cố (Mean Time To Detect)"),
    ("23", "MTTR", "Thời gian trung bình khắc phục sự cố (Mean Time To Repair)"),
    ("24", "OCC", "Kiểm soát tương tranh lạc quan (Optimistic Concurrency Control)"),
    ("25", "OEE", "Hiệu suất thiết bị tổng thể (Overall Equipment Effectiveness)"),
    ("26", "PMC", "Kế hoạch và điều độ vật tư sản xuất (Production Material Control)"),
    ("27", "PPH", "Năng suất sản lượng theo giờ trên mỗi lao động (Pieces Per Hour)"),
    ("28", "QC", "Quản lý và kiểm soát chất lượng (Quality Control)"),
    ("29", "R2", "Dịch vụ lưu trữ đối tượng không tính phí băng thông ra ngoài của Cloudflare"),
    ("30", "RBAC", "Hệ thống kiểm soát truy cập dựa trên vai trò (Role-Based Access Control)"),
    ("31", "UI/UX", "Giao diện người dùng và trải nghiệm người dùng (User Interface / User Experience)"),
]

for idx, (stt, term, desc) in enumerate(abbr_list, start=1):
    if idx < len(t10.rows):
        row = t10.rows[idx]
        set_cell_props(row.cells[0], stt, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_props(row.cells[1], term, bold=True, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_props(row.cells[2], desc, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

# -------------------------------------------------------------
# 3. LỜI CẢM ƠN (Paragraph 141)
# -------------------------------------------------------------
print("Step 3: Updating Lời cảm ơn...")
p141 = doc.paragraphs[141]
p141.text = (
    "Em xin trân trọng cảm ơn Ban Lãnh đạo Công Ty CPĐT Thái Bình Kiên Giang – TBS Group (Zone II) – "
    "SKECHERS R&D Center, Ban Tổng Giám đốc và người hướng dẫn TGĐ. Phạm Minh Tùng, cùng Phòng Hành Chính – "
    "Nhân Sự – Team Chuyển đổi số đã tạo mọi điều kiện thuận lợi nhất để em được thực tập, học hỏi và làm việc "
    "trong môi trường công nghiệp chuyên nghiệp, hiện đại. Các anh chị đã tận tình chỉ bảo, chia sẻ kinh nghiệm "
    "thực tế quý báu, từ quy trình điều hành sản xuất, tiêu chuẩn kiểm soát chất lượng quốc tế của thương hiệu "
    "SKECHERS đến việc hỗ trợ em hoàn thành xuất sắc các nhiệm vụ được giao, đặc biệt là trong quá trình phân tích, "
    "thiết kế, phát triển và triển khai hệ thống 'Quản Trị Vận Hành và Số Hóa Quy Trình Doanh Nghiệp trên nền tảng "
    "Serverless Edge – Tập đoàn TBS Group'."
)
p141.style = 'Normal'
for r in p141.runs:
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)

# -------------------------------------------------------------
# Exact Indices Definition
# -------------------------------------------------------------
MODAU = 162
TOMTAT = 172
C1 = 175
C2 = 201
C3 = 273
C4 = 287
C5 = 301
C6 = 311
C7 = 317
TL = 323

# -------------------------------------------------------------
# 4. PHẦN MỞ ĐẦU & TÓM TẮT ĐỒ ÁN
# -------------------------------------------------------------
print("Step 4: Populating Phần Mở Đầu & Tóm Tắt...")
format_p(doc.paragraphs[MODAU], "PHẦN MỞ ĐẦU", style='1', bold=True)

# 1. Lí do chọn đề tài
format_p(doc.paragraphs[MODAU + 1], "1. Lí do chọn đề tài", style='Heading 2', bold=True)
format_p(doc.paragraphs[MODAU + 2], 
    "Trong bối cảnh cách mạng công nghiệp 4.0 và xu hướng chuyển đổi số doanh nghiệp đang diễn ra mạnh mẽ, "
    "các tập đoàn sản xuất lớn phải đối mặt với bài toán tối ưu hóa quy trình điều hành và liên thông dữ liệu. "
    "Tại Tập đoàn TBS Group, đặc biệt là Cụm công nghiệp Kiên Giang (Zone II) kết hợp cùng Trung tâm Nghiên cứu & "
    "Phát triển mẫu SKECHERS R&D Center, khối lượng sản xuất hàng chục triệu đôi giày mỗi năm đòi hỏi sự phối hợp "
    "nhịp nhàng giữa hàng chục ngàn người lao động, các xưởng sản xuất vệ tinh và các khối phòng ban chức năng.\n\n"
    "Các hệ thống quản trị truyền thống vận hành trên máy chủ VPS tập trung thường bộc lộ nhiều điểm nghẽn: chi phí duy trì "
    "hạ tầng cao, độ trễ truyền tải lớn và nguy cơ gián đoạn dịch vụ khi mở rộng quy mô. Việc tiên phong ứng dụng công nghệ "
    "Serverless Edge Computing (Cloudflare Workers, D1, R2) mang lại giải pháp đột phá với độ trễ phản hồi cực thấp (< 50ms), "
    "khả năng tự động co giãn không giới hạn và loại bỏ chi phí quản trị máy chủ. Vì vậy, em đã chọn đề tài: 'Xây dựng hệ thống "
    "Quản Trị Vận Hành và Số Hóa Quy Trình Doanh Nghiệp trên nền tảng Serverless Edge – Tập đoàn TBS Group'."
)

# 2. Mục tiêu
format_p(doc.paragraphs[MODAU + 3], "2. Mục tiêu của báo cáo thực tập", style='Heading 2', bold=True)
format_p(doc.paragraphs[MODAU + 4],
    "Báo cáo tập trung vào 5 mục tiêu nghiên cứu và phát triển trọng tâm:\n"
    "1. Nghiên cứu kiến trúc Serverless Edge hiện đại trên nền tảng Cloudflare Workers, CSDL quan hệ D1 và bộ lưu trữ R2.\n"
    "2. Phân tích nghiệp vụ thực tế và xây dựng 10 phân hệ số hóa cốt lõi tại Văn phòng Chuỗi SKECHERS: My Tasks, Kanban Boards, "
    "Projects, Gemba Walk, CI-Kaizen, PPH Production Output, MMTB Maintenance, Quản trị 1-5-2, Nhân sự - Hành chính và Logistics.\n"
    "3. Thiết lập hệ thống bảo mật RBAC 7 tầng, loại bỏ triệt để các lỗ hổng OWASP API Security (BOLA, IDOR, Auth Bypass, PII).\n"
    "4. Tích hợp hệ thống thông báo đa kênh (Web Push VAPID, Zalo Notification Bot) và cơ chế tự động sao lưu dữ liệu sang Google Drive.\n"
    "5. Tiến hành kiểm thử tự động toàn diện (Pilot UAT 100 users, Multi-Role QA) và triển khai chính thức lên Production tại domain: "
    "https://vpchuoiskechers.tbsgroup2026.workers.dev."
)

# 3. Phạm vi
format_p(doc.paragraphs[MODAU + 5], "3. Phạm vi và cách thức thực tập", style='Heading 2', bold=True)
format_p(doc.paragraphs[MODAU + 6],
    "• Phạm vi thực hiện: Khối Văn phòng Chuỗi SKECHERS, Trung tâm R&D Center và cụm xưởng sản xuất Kiên Giang (Zone II). "
    "Người dùng trực tiếp gồm Ban Tổng Giám Đốc, Trưởng phòng ban (PMC, IE, QC, R&D, Nhân sự, Bảo trì) và 100 cán bộ công nhân viên Pilot.\n"
    "• Cách thức thực tập: Khảo sát hiện trường thực tế (Gemba Walk), phỏng vấn quy trình nghiệp vụ với các phòng ban, áp dụng mô hình "
    "Agile/Scrum trong phát triển phần mềm, xây dựng kịch bản kiểm thử tự động bằng Node.js và Puppeteer trước khi deploy lên Edge."
)

# 4. Kết quả đạt được
format_p(doc.paragraphs[MODAU + 7], "4. Kết quả đạt được", style='Heading 2', bold=True)
format_p(doc.paragraphs[MODAU + 8],
    "1. Hoàn thành hệ thống web application quy mô lớn gồm 472 file mã nguồn (144.560 dòng code TypeScript/React) và 1 Edge Worker dài 17.000 dòng code JavaScript.\n"
    "2. Vận hành CSDL Cloudflare D1 với 131 bảng dữ liệu, 19 migrations SQL và kết nối lưu trữ phi tập trung Cloudflare R2.\n"
    "3. Kiểm thử tự động đạt 100% PASS: 10/10 kịch bản Pilot UAT, 15/15 kịch bản phân quyền đa vai trò và vá hoàn tất 14 điểm rủi ro bảo mật.\n"
    "4. Triển khai chính thức lên mạng biên Cloudflare với thời gian khởi động worker chỉ 7ms, truy vấn D1 < 2ms, vận hành 24/7."
)

for pi in range(MODAU + 9, TOMTAT):
    doc.paragraphs[pi].text = ""

# Tóm tắt
format_p(doc.paragraphs[TOMTAT], "TÓM TẮT ĐỒ ÁN", style='1', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
format_p(doc.paragraphs[TOMTAT + 1],
    "Báo cáo thực tập tốt nghiệp trình bày quá trình nghiên cứu, thiết kế, phát triển và triển khai thực tế hệ thống 'Quản Trị "
    "Vận Hành và Số Hóa Quy Trình Doanh Nghiệp trên nền tảng Serverless Edge' tại Công Ty Cổ Phần Đầu Tư Thái Bình Kiên Giang – "
    "Tập đoàn TBS Group (Zone II) – SKECHERS R&D Center. Trước thách thức về độ trễ, chi phí hạ tầng và nguy cơ tắc nghẽn của hệ thống "
    "máy chủ VPS truyền thống trong bối cảnh sản xuất quy mô lớn, đề tài đã tiên phong ứng dụng kiến trúc điện toán biên không máy chủ "
    "(Serverless Edge Computing) dựa trên hệ sinh thái Cloudflare (Workers, D1 SQLite, R2 Storage) kết hợp cùng frontend Next.js 14 "
    "và Tailwind CSS. Hệ thống đã số hóa thành công 10 phân hệ nghiệp vụ cốt lõi: Quản lý công việc cá nhân (My Tasks), Bảng công việc "
    "phòng ban (Kanban), Dự án liên phòng ban và Cảnh báo điểm nghẽn (Projects & Bottleneck Detector), Kiểm soát hiện trường GEMBA Walk "
    "qua mã QR, Đăng ký và thẩm định sáng kiến cải tiến liên tục (CI-Kaizen), Quản lý sản lượng PPH nhà máy theo giờ, Bảo trì thiết bị "
    "MMTB (MTTD/MTTR, Pareto), Hệ thống quản trị 1-5-2 bảo mật mã PIN, Quản trị nhân sự định biên và Phân hệ Kế hoạch chuẩn bị và trung "
    "tâm phân phối. Đồng thời, hệ thống tích hợp sâu thông báo đẩy Web Push (VAPID), Zalo Notification Bot và cơ chế tự động sao lưu "
    "sang Google Drive qua Service Account RSA-2048. Toàn bộ nền tảng đã vượt qua các đợt kiểm thử tự động nghiêm ngặt (10/10 Pilot UAT, "
    "15/15 Multi-Role QA, 14/14 Security Audit) và chính thức được đưa vào vận hành tại https://vpchuoiskechers.tbsgroup2026.workers.dev, "
    "đáp ứng xuất sắc các tiêu chuẩn vận hành công nghiệp khắt khe của thương hiệu SKECHERS toàn cầu.\n\n"
    "Từ khóa: Serverless Edge Computing, Cloudflare Workers, Cloudflare D1, Quản Trị Vận Hành, Chuyển Đổi Số, TBS Group, SKECHERS R&D."
)

for pi in range(TOMTAT + 2, C1):
    doc.paragraphs[pi].text = ""

# -------------------------------------------------------------
# 5. CHƯƠNG 1: TỔNG QUAN VỀ CÔNG TY
# -------------------------------------------------------------
print("Step 5: Populating Chương 1...")
format_p(doc.paragraphs[C1], "CHƯƠNG 1: TỔNG QUAN VỀ CÔNG TY CPĐT THÁI BÌNH KIÊN GIANG – TBS GROUP (ZONE II) – SKECHERS R&D CENTER", style='1', bold=True)

# 1.1 Thông tin đơn vị
format_p(doc.paragraphs[C1 + 1], "1.1. Thông tin đơn vị", style='2', bold=True)
p_c1_1_c = doc.paragraphs[C1 + 2]
p_c1_1_c.text = ""
p_c1_1_c.style = 'Normal'
add_highlighted_run(p_c1_1_c, "Tập đoàn Thái Bình Shoes (TBS Group) là một trong những tập đoàn công nghiệp đa ngành hàng đầu tại Việt Nam, đặc biệt khẳng định vị thế dẫn đầu trong lĩnh vực sản xuất giày da và gia công túi xách xuất khẩu cho các thương hiệu danh tiếng toàn cầu. Trong chiến lược mở rộng chuỗi cung ứng sản xuất tại khu vực phía Nam, Công Ty Cổ Phần Đầu Tư Thái Bình Kiên Giang (Zone II) được thành lập và đầu tư trọng điểm, trở thành trung tâm nghiên cứu phát triển mẫu và sản xuất mũi nhọn của thương hiệu SKECHERS.\n\nThông tin cơ bản của đơn vị thực tập:\n• Tên đơn vị: Công Ty Cổ Phần Đầu Tư Thái Bình Kiên Giang – Tập đoàn TBS Group (Zone II) – SKECHERS R&D Center\n• Tên giao dịch quốc tế: TBS Kiên Giang Investment Joint Stock Company – SKECHERS Division\n• Mã số thuế: ")
add_highlighted_run(p_c1_1_c, "[CẦN BỔ SUNG: Mã số thuế Công ty CPĐT Thái Bình Kiên Giang]", is_highlight=True)
add_highlighted_run(p_c1_1_c, "\n• Địa chỉ trụ sở: Khu công nghiệp Thạnh Lộc, xã Thạnh Lộc, huyện Châu Thành, tỉnh Kiên Giang (Zone II - TBS Group)\n• Điện thoại văn phòng: ")
add_highlighted_run(p_c1_1_c, "[CẦN BỔ SUNG: Số điện thoại văn phòng]", is_highlight=True)
add_highlighted_run(p_c1_1_c, "\n• Ngày thành lập chi nhánh: ")
add_highlighted_run(p_c1_1_c, "[CẦN BỔ SUNG: Ngày cấp phép thành lập chi nhánh Kiên Giang]", is_highlight=True)
add_highlighted_run(p_c1_1_c, " (Tập đoàn TBS Group thành lập từ năm 1989)\n• Website tập đoàn: https://tbsgroup.vn\n• Lãnh đạo phụ trách đơn vị: Tổng Giám Đốc Phạm Minh Tùng\n• Chủ tịch Hội đồng Quản trị Tập đoàn: Ông Nguyễn Đức Thuấn")

def find_local(query, start, end):
    for i in range(start, end):
        if query.lower() in doc.paragraphs[i].text.lower():
            return i
    return -1

idx_1_2 = find_local("Quá trình thành lập", C1, C2)
if idx_1_2 > -1:
    for pi in range(C1 + 3, idx_1_2):
        doc.paragraphs[pi].text = ""
    format_p(doc.paragraphs[idx_1_2], "1.2. Quá trình thành lập và phát triển", style='2', bold=True)
    format_p(doc.paragraphs[idx_1_2 + 1],
        "Hành trình phát triển của Tập đoàn TBS Group trải qua hơn 35 năm bền bỉ, khởi đầu từ một xưởng sản xuất giày da thủ công nhỏ "
        "thành lập năm 1989 tại tỉnh Bình Dương. Với tầm nhìn chiến lược của Chủ tịch HĐQT Nguyễn Đức Thuấn, tập đoàn đã nhanh chóng vươn lên, "
        "trở thành đối tác sản xuất chiến lược toàn cầu của các tập đoàn giày thể thao và thời trang hàng đầu thế giới như SKECHERS, Decathlon, Coach.\n\n"
        "Giai đoạn những năm 2010 đến nay, nhằm đón đầu làn sóng dịch chuyển sản xuất và tối ưu hóa nguồn lực lao động tại Đồng bằng sông Cửu Long, "
        "TBS Group đã thành lập Công Ty CPĐT Thái Bình Kiên Giang (Zone II). Đơn vị nhanh chóng được quy hoạch hiện đại với cụm xưởng sản xuất "
        "quy mô lớn và Trung tâm Nghiên cứu & Phát triển mẫu (SKECHERS R&D Center). Năm 2026, TBS Kiên Giang (Zone II) đóng vai trò là hạt nhân "
        "tiên phong trong chiến lược Chuyển đổi số 4.0 của toàn tập đoàn, đẩy mạnh tự động hóa chuyền may, số hóa chỉ số OEE/PPH theo thời gian thực "
        "và áp dụng công nghệ điện toán biên vào công tác quản trị điều hành sản xuất."
    )

idx_1_3 = find_local("Chức năng nhiệm vụ", C1, C2)
if idx_1_3 > -1:
    format_p(doc.paragraphs[idx_1_3], "1.3. Chức năng và nhiệm vụ", style='2', bold=True)

idx_1_3_1 = find_local("1.3.1", C1, C2)
if idx_1_3_1 > -1:
    format_p(doc.paragraphs[idx_1_3_1], "1.3.1. Chức năng của đơn vị", style='Heading 3', bold=True)
    format_p(doc.paragraphs[idx_1_3_1 + 1],
        "1. Đầu não điều phối chuỗi cung ứng sản xuất SKECHERS: Quản lý kế hoạch cung ứng nguyên vật liệu, theo dõi đơn hàng xuất khẩu "
        "và điều phối tiến độ gia công giữa các phân xưởng trong toàn bộ Zone II.\n"
        "2. Trung tâm Nghiên cứu & Phát triển công nghệ (R&D Center): Tiếp nhận bản vẽ thiết kế 3D từ nhãn hàng SKECHERS, nghiên cứu phom dáng, "
        "chế thử mẫu giày thể thao, thử nghiệm độ bền vật liệu tại phòng Lab và chuyển giao quy trình công nghệ cho các nhà máy sản xuất hàng loạt.\n"
        "3. Quản trị chất lượng theo tiêu chuẩn toàn cầu: Kiểm định nghiêm ngặt chất lượng sản phẩm qua từng công đoạn (Chặt, May, Gò, Đóng gói) "
        "đảm bảo chuẩn chất lượng AQL 1.0 của thương hiệu SKECHERS."
    )

idx_1_3_2 = find_local("1.3.2", C1, C2)
if idx_1_3_2 > -1:
    format_p(doc.paragraphs[idx_1_3_2], "1.3.2. Nhiệm vụ của đơn vị", style='Heading 3', bold=True)
    format_p(doc.paragraphs[idx_1_3_2 + 1],
        "1. Hoàn thành vượt mức kế hoạch sản lượng xuất khẩu được Tập đoàn TBS Group và đối tác SKECHERS giao phó hàng năm.\n"
        "2. Thực hiện nghiêm ngặt các quy định về an toàn lao động, kiểm soát hiện trường sản xuất (Gemba Walk), tiêu chuẩn 5S và phòng cháy chữa cháy.\n"
        "3. Quản lý, bảo trì và nâng cao hiệu suất hoạt động của hệ thống hàng ngàn máy móc thiết bị (MMTB) hiện đại.\n"
        "4. Đẩy mạnh phong trào Cải tiến liên tục (CI-Kaizen), khuyến khích CBCNV đóng góp sáng kiến nâng cao năng suất và tiết kiệm chi phí.\n"
        "5. Xây dựng môi trường làm việc nhân văn, chuyên nghiệp, chăm lo đời sống vật chất và tinh thần cho hàng ngàn người lao động tại địa phương."
    )

idx_1_4 = find_local("Sơ đồ tổ chức", C1, C2)
if idx_1_4 > -1:
    format_p(doc.paragraphs[idx_1_4], "1.4. Sơ đồ cơ cấu tổ chức bộ máy quản lý", style='2', bold=True)
    p_sd_c = doc.paragraphs[idx_1_4 + 1]
    p_sd_c.text = ""
    add_highlighted_run(p_sd_c, "Cơ cấu tổ chức của Công Ty CPĐT Thái Bình Kiên Giang – TBS Group (Zone II) – SKECHERS R&D Center được vận hành theo mô hình quản trị ma trận liên thông, kết hợp chặt chẽ giữa quản trị chức năng và chuỗi giá trị sản phẩm:\n• Hội đồng Quản trị & Ban Tổng Giám Đốc: Định hướng chiến lược phát triển, phê duyệt kế hoạch tài chính và quyết định đầu tư công nghệ. Đứng đầu là Chủ tịch HĐQT Nguyễn Đức Thuấn và Tổng Giám Đốc Phạm Minh Tùng.\n• Khối Văn phòng Chuỗi SKECHERS: Đầu mối điều hành chuỗi cung ứng, phụ trách kế hoạch vật tư (PMC) và trung tâm phân phối.\n• Trung tâm Nghiên cứu & Phát triển mẫu (SKECHERS R&D Center): Phụ trách kỹ thuật 3D CAD/CAM, chế thử mẫu và thử nghiệm phòng LAB.\n• Khối Quản lý Chất lượng (QA/QC): Giám sát độc lập tiêu chuẩn chất lượng sản phẩm từ khâu nguyên phụ liệu đến thành phẩm.\n• Khối Cải tiến liên tục (CI) & Kỹ thuật công nghiệp (IE): Tối ưu hóa layout chuyền may, định mức thời gian và tổ chức phong trào Kaizen.\n• Khối Quản trị Nguồn nhân lực & Hành chính: Tuyển dụng, đào tạo, chính sách tiền lương (C&B) và quản lý cơ sở vật chất.\n• Khối Sản xuất & Hệ thống Phân xưởng: Trực tiếp điều hành các xưởng may, xưởng chặt, xưởng gò đế và hoàn thiện sản phẩm.\n\nSơ đồ trực quan: ")
    add_highlighted_run(p_sd_c, "[CẦN BỔ SUNG: Sơ đồ cơ cấu tổ chức bản vẽ chính thức của đơn vị]", is_highlight=True)
    if idx_1_4 + 2 < len(doc.paragraphs):
        add_caption(doc.paragraphs[idx_1_4 + 2], "Hình 1.1: Sơ đồ cơ cấu tổ chức Công Ty CPĐT Thái Bình Kiên Giang – TBS Group (Zone II)")

idx_1_5 = find_local("Nhiệm vụ được giao", C1, C2)
if idx_1_5 > -1:
    format_p(doc.paragraphs[idx_1_5], "1.5. Nhiệm vụ được giao trong quá trình thực tập", style='2', bold=True)
    format_p(doc.paragraphs[idx_1_5 + 1],
        "Trong đợt thực tập tốt nghiệp, em vinh dự được nhận vào làm việc tại Team Chuyển đổi số trực thuộc Ban Tổng Giám Đốc và "
        "Phòng Hành chính - Nhân sự. Dưới sự chỉ đạo trực tiếp của TGĐ. Phạm Minh Tùng, em được phân công đảm nhiệm các nhiệm vụ trọng tâm:\n"
        "1. Khảo sát thực địa và đánh giá mức độ số hóa của các quy trình nghiệp vụ hiện hành tại các phòng ban thuộc Zone II.\n"
        "2. Trực tiếp chịu trách nhiệm kiến trúc và phát triển phần mềm toàn trình (Fullstack): xây dựng frontend giao diện với Next.js 14, "
        "lập trình backend Serverless Edge Worker với Cloudflare Workers, thiết kế cấu trúc CSDL Cloudflare D1 và bộ nhớ lưu trữ Cloudflare R2.\n"
        "3. Thiết kế hệ thống phân quyền đa vai trò (RBAC), khắc phục triệt để các lỗ hổng bảo mật và xây dựng cơ chế xác thực JWT an toàn.\n"
        "4. Tích hợp các dịch vụ bên ngoài: hệ thống thông báo đẩy Web Push VAPID, chatbot thông báo Zalo qua Webhook và tự động sao lưu dữ liệu sang Google Drive.\n"
        "5. Tổ chức kiểm thử tự động toàn diện, chủ trì triển khai chương trình thử nghiệm Pilot 100 người dùng và đóng gói triển khai chính thức hệ thống lên Production."
    )

idx_1_6 = find_local("Mục tiêu và kế hoạch thực tập", C1, C2)
if idx_1_6 > -1:
    format_p(doc.paragraphs[idx_1_6], "1.6. Mục tiêu và kế hoạch thực tập", style='2', bold=True)
    format_p(doc.paragraphs[idx_1_6 + 1],
        "Mục tiêu xuyên suốt quá trình thực tập là rèn luyện tác phong làm việc công nghiệp, chuyển hóa các kiến thức học thuật tại trường "
        "thành giải pháp công nghệ có khả năng giải quyết trực tiếp các bài toán thực tiễn phức tạp của doanh nghiệp. "
        "Kế hoạch thực tập được phân bổ khoa học qua 10 tuần, từ ngày 03/8/2026 đến ngày 10/10/2026, tuần tự từ khâu khảo sát phân tích, "
        "thiết kế kiến trúc, phát triển các phân hệ cốt lõi, tích hợp bảo mật, kiểm thử tự động đến đóng gói triển khai sản phẩm."
    )

idx_1_7 = find_local("Kiến thức và kỹ năng đã áp dụng", C1, C2)
if idx_1_7 > -1:
    format_p(doc.paragraphs[idx_1_7], "1.7. Kiến thức và kỹ năng đã áp dụng", style='2', bold=True)
    format_p(doc.paragraphs[idx_1_7 + 1],
        "Quá trình thực hiện đề tài đã giúp em vận dụng và phát huy tối đa các khối kiến thức và kỹ năng chuyên ngành:\n"
        "• Kiến thức lý thuyết: Kỹ thuật lập trình hiện đại (TypeScript/React), kiến trúc hệ thống phân tán, cơ sở dữ liệu quan hệ (SQL/SQLite), "
        "an toàn thông tin web (OWASP Top 10 API Security), mạng máy tính và giao thức HTTP/SSE/WebSocket.\n"
        "• Kỹ năng công nghệ mới: Làm chủ hệ sinh thái Serverless Edge Computing của Cloudflare (Workers V8 Isolate, D1 Database, R2 Object Storage), "
        "lập trình tối ưu hóa giao diện với Tailwind CSS, xử lý dữ liệu lớn với XLSX và Chart.js/Recharts.\n"
        "• Kỹ năng mềm: Giao tiếp hiệu quả với cán bộ công nhân viên và lãnh đạo cấp cao, kỹ năng giải quyết vấn đề (troubleshooting), "
        "kỹ năng quản lý thời gian và viết tài liệu kỹ thuật chuẩn mực."
    )

for pi in range(idx_1_7 + 2, C2):
    doc.paragraphs[pi].text = ""

# -------------------------------------------------------------
# 6. CHƯƠNG 2, 3, 4, 5, 6, 7 & TÀI LIỆU THAM KHẢO
# -------------------------------------------------------------
print("Step 6: Populating Chapters 2 to 7...")

import populate_chapters_2_to_7 as c27

format_p(doc.paragraphs[C2], "CHƯƠNG 2: TỔNG QUAN VỀ HỆ THỐNG QUẢN TRỊ VẬN HÀNH VÀ SỐ HÓA TBS GROUP (ZONE II)", style='1', bold=True)
format_p(doc.paragraphs[C2 + 1], c27.c2_content, style='Normal')
for pi in range(C2 + 2, C3):
    doc.paragraphs[pi].text = ""

format_p(doc.paragraphs[C3], "CHƯƠNG 3: THIẾT KẾ HỆ THỐNG", style='1', bold=True)
format_p(doc.paragraphs[C3 + 1], c27.c3_content, style='Normal')
for pi in range(C3 + 2, C4):
    doc.paragraphs[pi].text = ""

format_p(doc.paragraphs[C4], "CHƯƠNG 4: TRIỂN KHAI HỆ THỐNG", style='1', bold=True)
format_p(doc.paragraphs[C4 + 1], c27.c4_content, style='Normal')

# Add real screenshots into Chapter 4
p_img_container = doc.paragraphs[C4 + 2]
p_img_container.text = "\nDanh mục hình ảnh giao diện thực tế hệ thống đã triển khai tại TBS Group (Zone II):\n"
p_img_container.style = 'Normal'

screenshots = [
    ("docs/screenshots/01_work_main_dashboard.png", "Hình 4.1: Bảng điều khiển trung tâm /work trên máy tính để bàn (Desktop Dashboard)"),
    ("docs/screenshots/01_work_main_dashboard_mobile.png", "Hình 4.2: Giao diện Bảng điều khiển tối ưu hóa trên điện thoại di động (Mobile Dashboard 390px)"),
    ("docs/screenshots/03_work_kaizen_dashboard.png", "Hình 4.3: Giao diện Phân hệ Sáng kiến cải tiến liên tục CI-Kaizen & Bảng xếp hạng"),
    ("docs/screenshots/04_work_gemba.png", "Hình 4.4: Giao diện Phân hệ Kiểm soát hiện trường sản xuất GEMBA Walk & Quét mã QR"),
    ("docs/screenshots/06_maintenance_kpi.png", "Hình 4.5: Giao diện Phân hệ Quản trị máy móc thiết bị và bảo trì MMTB (Chỉ số 6 nhà máy)"),
    ("docs/screenshots/07_pph_view.png", "Hình 4.6: Giao diện Phân hệ Quản lý sản lượng PPH nhà máy theo giờ & Màn hình TV xưởng"),
    ("docs/screenshots/rename_module/live_prod_logistics_route.png", "Hình 4.7: Giao diện Phân hệ Kế hoạch chuẩn bị và trung tâm phân phối trên Production"),
]

for img_path, caption in screenshots:
    if os.path.exists(img_path):
        p_pic = doc.add_paragraph()
        p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pic.paragraph_format.space_before = Pt(6)
        p_pic.paragraph_format.space_after = Pt(2)
        try:
            p_pic.add_run().add_picture(img_path, width=Inches(5.6))
            p_cap = doc.add_paragraph()
            add_caption(p_cap, caption)
            p_img_container._element.addnext(p_cap._element)
            p_img_container._element.addnext(p_pic._element)
        except Exception as e:
            print(f"Warning inserting image {img_path}: {e}")

for pi in range(C4 + 3, C5):
    doc.paragraphs[pi].text = ""

format_p(doc.paragraphs[C5], "CHƯƠNG 5: KIỂM THỬ VÀ ĐÁNH GIÁ HỆ THỐNG", style='1', bold=True)
format_p(doc.paragraphs[C5 + 1], c27.c5_content, style='Normal')
for pi in range(C5 + 2, C6):
    doc.paragraphs[pi].text = ""

format_p(doc.paragraphs[C6], "CHƯƠNG 6: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", style='1', bold=True)
format_p(doc.paragraphs[C6 + 1], c27.c6_content, style='Normal')
for pi in range(C6 + 2, C7):
    doc.paragraphs[pi].text = ""

format_p(doc.paragraphs[C7], "CHƯƠNG 7: ĐÁNH GIÁ THỰC TIỄN VÀ MỘT SỐ KIẾN NGHỊ", style='1', bold=True)
format_p(doc.paragraphs[C7 + 1], c27.c7_content, style='Normal')
for pi in range(C7 + 2, TL):
    doc.paragraphs[pi].text = ""

format_p(doc.paragraphs[TL], "TÀI LIỆU THAM KHẢO", style='1', bold=True)
if TL + 1 < len(doc.paragraphs):
    p_tl_c = doc.paragraphs[TL + 1]
else:
    p_tl_c = doc.add_paragraph()
format_p(p_tl_c, c27.references_content, style='Normal')
for pi in range(TL + 2, len(doc.paragraphs)):
    doc.paragraphs[pi].text = ""

# -------------------------------------------------------------
# 7. SAVE AND VERIFY XML PARTS
# -------------------------------------------------------------
print("Saving populated document to:", DOCX_OUT)
doc.save(DOCX_OUT)

print("Verifying XML parts for forbidden terms...")
forbidden = ['việc làm xanh', 'viec lam xanh', 'nhân kiệt', 'nhan kiet', 'cung ứng nhân lực', 'cung ung nhan luc']
found_any = False

with zipfile.ZipFile(DOCX_OUT, 'r') as z:
    for filename in z.namelist():
        if filename.endswith(('.xml', '.rels')):
            content = z.read(filename).decode('utf-8', errors='ignore')
            content_lower = content.lower()
            for f in forbidden:
                if f in content_lower:
                    cnt = content_lower.count(f)
                    print(f"❌ [ALERT] In {filename}: '{f}' found {cnt} times!")
                    found_any = True

if not found_any:
    print("🎉 100% CLEAN: 0 forbidden terms found in all XML parts!")

print("=== FINALIZATION COMPLETE ===")
