# -*- coding: utf-8 -*-
"""
Full Document Builder for Internship Report (Bao_cao_thuc_tap_tot_nghiep.docx)
Author: Antigravity AI
Target: TBS Group - Zone II, Chuỗi Skechers
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from report_core import (
    set_cell_background, set_cell_margins, set_table_borders,
    add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_bullet_p, add_figure, add_table_title
)

DOC_PATH = r"D:\Work\TBS II\Bao_cao_thuc_tap_tot_nghiep.docx"

def set_page_margins(section):
    # A4 standard
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    # Top 2 cm, Bottom 2 cm, Left 3 cm, Right 2 cm
    section.top_margin = Inches(2.0 / 2.54)
    section.bottom_margin = Inches(2.0 / 2.54)
    section.left_margin = Inches(3.0 / 2.54)
    section.right_margin = Inches(2.0 / 2.54)

def add_page_number_field(p):
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p._element.append(fldSimple)

def setup_footer(section, is_roman=False, start_num=1):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    
    # Configure page numbering XML
    sectPr = section._sectPr
    # Remove existing pgNumType if any
    for el in sectPr.findall(qn('w:pgNumType')):
        sectPr.remove(el)
        
    pgNumType = OxmlElement('w:pgNumType')
    if is_roman:
        pgNumType.set(qn('w:fmt'), 'lowerRoman')
    else:
        pgNumType.set(qn('w:fmt'), 'decimal')
    pgNumType.set(qn('w:start'), str(start_num))
    sectPr.append(pgNumType)
    
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 116, 139)
    add_page_number_field(p)

def build_cover_page(doc, is_sub_cover=False):
    # Title box / Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(2)
    r1 = p_header.add_run("UBND THÀNH PHỐ HỒ CHÍ MINH\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(13)
    r1.font.bold = True
    
    r2 = p_header.add_run("TRƯỜNG ĐẠI HỌC THỦ DẦU MỘT\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 104, 56) # Emerald Green
    
    r3 = p_header.add_run("VIỆN CÔNG NGHỆ SỐ\n")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(13)
    r3.font.bold = True
    
    r_line = p_header.add_run("-----------------------***-----------------------\n")
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(11)
    r_line.font.color.rgb = RGBColor(148, 163, 184)
    
    # Large Space
    for _ in range(3):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    
    rt1 = p_title.add_run("BÁO CÁO CHUYÊN ĐỀ\n")
    rt1.font.name = 'Times New Roman'
    rt1.font.size = Pt(16)
    rt1.font.bold = True
    
    rt2 = p_title.add_run("THỰC TẬP TỐT NGHIỆP\n\n")
    rt2.font.name = 'Times New Roman'
    rt2.font.size = Pt(20)
    rt2.font.bold = True
    rt2.font.color.rgb = RGBColor(0, 104, 56)
    
    rt_work = p_title.add_run("Nơi thực tập:\nTẬP ĐOÀN TBS GROUP – ZONE II, CHUỖI SKECHERS\n(CÔNG TY CP ĐẦU TƯ THÁI BÌNH – NHÀ MÁY 2)\n\n")
    rt_work.font.name = 'Times New Roman'
    rt_work.font.size = Pt(13.5)
    rt_work.font.bold = True
    
    rt_topic_label = p_title.add_run("TÊN ĐỀ TÀI:\n")
    rt_topic_label.font.name = 'Times New Roman'
    rt_topic_label.font.size = Pt(13)
    rt_topic_label.font.bold = True
    
    # Requirement: Leave empty dotted line
    rt_topic_dot = p_title.add_run("………………………………………………………………………………………………………\n………………………………………………………………………………………………………\n")
    rt_topic_dot.font.name = 'Times New Roman'
    rt_topic_dot.font.size = Pt(13)
    rt_topic_dot.font.bold = True
    
    for _ in range(4):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(0)
        
    # Info block (Table 2 columns)
    tbl = doc.add_table(rows=7, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("CÁN BỘ HƯỚNG DẪN TẠI ĐƠN VỊ:", "Dư Thị Thanh Tình"),
        ("GIẢNG VIÊN HƯỚNG DẪN:", "ThS. Trần Bá Minh Sơn"),
        ("SINH VIÊN THỰC HIỆN:", "[CẦN BỔ SUNG: Họ và tên sinh viên]"),
        ("MÃ SỐ SINH VIÊN:", "[CẦN BỔ SUNG: MSSV]"),
        ("CHUYÊN NGÀNH:", "Công Nghệ Thông Tin"),
        ("LỚP:", "[CẦN BỔ SUNG: Tên lớp]"),
        ("NIÊN KHÓA:", "[CẦN BỔ SUNG: Niên khóa]"),
    ]
    
    for i, (label, val) in enumerate(info_data):
        cell_lbl = tbl.cell(i, 0)
        cell_val = tbl.cell(i, 1)
        
        cell_lbl.width = Inches(2.7)
        cell_val.width = Inches(3.4)
        
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(2)
        p_lbl.paragraph_format.space_after = Pt(2)
        r_l = p_lbl.add_run(label)
        r_l.font.name = 'Times New Roman'
        r_l.font.size = Pt(12.5)
        r_l.font.bold = True
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(2)
        p_val.paragraph_format.space_after = Pt(2)
        r_v = p_val.add_run(val)
        r_v.font.name = 'Times New Roman'
        r_v.font.size = Pt(12.5)
        if "[" in val:
            r_v.font.italic = True
            r_v.font.color.rgb = RGBColor(194, 65, 12) # Amber highlight
        else:
            r_v.font.bold = True
            
        set_cell_margins(cell_lbl, top=40, bottom=40, left=50, right=50)
        set_cell_margins(cell_val, top=40, bottom=40, left=50, right=50)
        
    for _ in range(3):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(0)
        
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run("TP. Hồ Chí Minh, năm 2026")
    r_loc.font.name = 'Times New Roman'
    r_loc.font.size = Pt(12.5)
    r_loc.font.bold = True

print("Loaded Document Builder part 1.")
