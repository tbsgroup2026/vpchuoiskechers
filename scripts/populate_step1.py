import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os
import sys
import zipfile
import re

sys.stdout.reconfigure(encoding='utf-8')

DOCX_IN = '2224802010738_PhamNguyenAnhHuy.docx'
DOCX_OUT = '2224802010738_PhamNguyenAnhHuy.docx'

print("=== STARTING REPORT POPULATION ENGINE ===")

doc = Document(DOCX_IN)

# ---------------------------------------------------------
# Helper functions
# ---------------------------------------------------------
def set_cell_text(cell, text, bold=False, font_size=11, font_name="Times New Roman", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_table_after_paragraph(p, rows, cols, col_widths=None):
    tbl = p._element.getparent().add_next(OxmlElement('w:tbl'))
    # Use python-docx Table wrapper
    # Alternatively, add table via doc.add_table and move its xml
    table = doc.add_table(rows=rows, cols=cols)
    p._element.addnext(table._element)
    return table

# ---------------------------------------------------------
# 1. UPDATE TABLE 2 (10-WEEK PLAN)
# ---------------------------------------------------------
print("Updating Table 2 (10-Week Plan)...")
t2 = doc.tables[2]

weeks_data = [
    # Tuần, Ngày, Nội dung
    ("1", "03/8 – 08/8", "Tìm hiểu và phân tích tổng quan dự án"),
    ("2", "10/8 – 15/8", "Xây dựng và phát triển các chức năng hệ thống"),
    ("3", "17/8 – 22/8", "Xây dựng phân hệ Quản lý công việc cá nhân (My Tasks), Bảng công việc phòng ban (Kanban) và Dự án liên phòng ban (Projects)"),
    ("4", "24/8 – 29/8", "Xây dựng phân hệ Kiểm soát hiện trường GEMBA Walk và Đăng ký sáng kiến cải tiến liên tục (CI-Kaizen)"),
    ("5", "31/8 – 05/9", "Xây dựng phân hệ Quản lý sản lượng PPH nhà máy và tích hợp dữ liệu Máy móc thiết bị (MMTB)"),
    ("6", "07/9 – 12/9", "Tích hợp hệ thống thông báo đa kênh: Web Push (VAPID), Zalo Notification Service và Service Worker nền tảng"),
    ("7", "14/9 – 19/9", "Hoàn thiện hệ thống phân quyền đa cấp (RBAC), bảo mật xác thực (JWT/Session) và cơ chế Audit Log / Idempotency"),
    ("8", "21/9 – 26/9", "Xây dựng cơ chế sao lưu tự động (Incremental Backup sang Google Drive), quản lý file media Cloudflare R2 và tối ưu nén ảnh client"),
    ("9", "28/9 – 03/10", "Kiểm thử toàn diện hệ thống: UAT 100 người dùng, kiểm thử bảo mật đa vai trò, kiểm thử API và sửa lỗi hồi quy"),
    ("10", "05/10 – 10/10", "Triển khai hệ thống lên Cloudflare Workers (Production), đánh giá hiệu năng và hoàn thiện báo cáo thực tập"),
]

for idx, (w_num, w_date, w_content) in enumerate(weeks_data, start=1):
    row = t2.rows[idx]
    # Col 0: Tuần lễ
    set_cell_text(row.cells[0], w_num, bold=True, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    # Col 1: Từ ngày ... đến ...
    set_cell_text(row.cells[1], w_date, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    # Col 2: In case merged with Col 1
    if len(row.cells) > 2 and row.cells[2] != row.cells[1]:
        set_cell_text(row.cells[2], w_date, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    # Col 3: Nội dung
    content_cell = row.cells[3] if len(row.cells) > 3 else row.cells[-2]
    set_cell_text(content_cell, w_content, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    # Col 4: Ghi chú
    note_cell = row.cells[4] if len(row.cells) > 4 else row.cells[-1]
    set_cell_text(note_cell, "", bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

print("✅ Table 2 updated successfully.")

# ---------------------------------------------------------
# 2. UPDATE TABLE 10 (ABBREVIATIONS - 31 TERMS)
# ---------------------------------------------------------
print("Updating Table 10 (Abbreviations)...")
t10 = doc.tables[10]

abbr_data = [
    ("1", "API", "Giao diện lập trình ứng dụng (Application Programming Interface)"),
    ("2", "AQL", "Giới hạn chất lượng chấp nhận (Acceptance Quality Limit)"),
    ("3", "BOLA", "Lỗ hổng kiểm soát truy cập cấp đối tượng (Broken Object Level Authorization)"),
    ("4", "CBCNV", "Cán bộ công nhân viên"),
    ("5", "CI", "Cải tiến liên tục (Continuous Improvement)"),
    ("6", "CI/CD", "Tích hợp liên tục và triển khai liên tục (Continuous Integration / Continuous Deployment)"),
    ("7", "CSP", "Chính sách bảo mật nội dung (Content Security Policy)"),
    ("8", "CSS", "Ngôn ngữ định kiểu tầng (Cascading Style Sheets)"),
    ("9", "D1", "Cơ sở dữ liệu quan hệ SQL phân tán tại mạng biên của Cloudflare"),
    ("10", "DOM", "Mô hình đối tượng tài liệu (Document Object Model)"),
    ("11", "ERD", "Sơ đồ quan hệ thực thể (Entity Relationship Diagram)"),
    ("12", "F5", "Thao tác làm mới tải lại trang trên trình duyệt web"),
    ("13", "GEMBA", "Hiện trường sản xuất thực tế theo triết lý Lean Manufacturing"),
    ("14", "HĐQT", "Hội đồng quản trị"),
    ("15", "HTML", "Ngôn ngữ đánh dấu siêu văn bản (HyperText Markup Language)"),
    ("16", "HTTP", "Giao thức truyền tải siêu văn bản (HyperText Transfer Protocol)"),
    ("17", "IDOR", "Tham chiếu đối tượng trực tiếp không an toàn (Insecure Direct Object References)"),
    ("18", "IE", "Kỹ thuật công nghiệp và định mức lao động (Industrial Engineering)"),
    ("19", "JWT", "Chuỗi mã định danh phiên an toàn (JSON Web Token)"),
    ("20", "KPI", "Chỉ số đánh giá hiệu quả công việc trọng yếu (Key Performance Indicator)"),
    ("21", "MMTB", "Máy móc thiết bị"),
    ("22", "MTTD", "Thời gian trung bình phát hiện sự cố (Mean Time To Detect)"),
    ("23", "MTTR", "Thời gian trung bình khắc phục sự cố (Mean Time To Repair)"),
    ("24", "OCC", "Kiểm soát tương tranh lạc quan (Optimistic Concurrency Control)"),
    ("25", "OEE", "Hiệu suất thiết bị tổng thể (Overall Equipment Effectiveness)"),
    ("26", "PMC", "Kế hoạch và điều độ vật tư sản xuất (Production Material Control)"),
    ("27", "PPH", "Năng suất sản lượng theo giờ trên mỗi lao động (Pieces Per Hour)"),
    ("28", "QC", "Quản lý và kiểm soát chất lượng (Quality Control)"),
    ("29", "R2", "Dịch vụ lưu trữ đối tượng không tính phí băng thông ra ngoài của Cloudflare"),
    ("30", "RBAC", "Hệ thống kiểm soát truy cập dựa trên vai trò (Role-Based Access Control)"),
    ("31", "UI/UX", "Giao diện người dùng và trải nghiệm người dùng (User Interface / User Experience)"),
]

for idx, (stt, term, desc) in enumerate(abbr_data, start=1):
    if idx < len(t10.rows):
        row = t10.rows[idx]
        set_cell_text(row.cells[0], stt, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], term, bold=True, font_size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[2], desc, bold=False, font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

print("✅ Table 10 updated successfully.")

# ---------------------------------------------------------
# 3. UPDATE LỜI CẢM ƠN (Paragraph 141)
# ---------------------------------------------------------
print("Updating Lời cảm ơn...")
p141 = doc.paragraphs[141]
p141.text = (
    "Em xin trân trọng cảm ơn Ban Lãnh đạo Công Ty CPĐT Thái Bình Kiên Giang – TBS Group (Zone II) – "
    "SKECHERS R&D Center, Ban Tổng Giám đốc và người hướng dẫn TGĐ. Phạm Minh Tùng, cùng Phòng Hành Chính – "
    "Nhân Sự – Team Chuyển đổi số đã tạo mọi điều kiện thuận lợi nhất để em được thực tập, học hỏi và làm việc "
    "trong môi trường công nghiệp chuyên nghiệp, hiện đại. Các anh chị đã tận tình chỉ bảo, chia sẻ kinh nghiệm "
    "thực tế quý báu, từ quy trình điều hành sản xuất, tiêu chuẩn kiểm soát chất lượng quốc tế của thương hiệu "
    "SKECHERS đến việc hỗ trợ em hoàn thành xuất sắc các nhiệm vụ được giao, đặc biệt là trong quá trình phân tích, "
    "thiết kế, phát triển và triển khai hệ thống 'Quản Trị Vận Hành và Số Hóa Quy Trình Doanh Nghiệp trên nền tảng "
    "Serverless Edge – Tập đoàn TBS Group'."
)
p141.style = 'Normal'
for r in p141.runs:
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.italic = False

print("✅ Lời cảm ơn updated.")

doc.save("scratch_step1.docx")
print("Saved checkpoint to scratch_step1.docx")
